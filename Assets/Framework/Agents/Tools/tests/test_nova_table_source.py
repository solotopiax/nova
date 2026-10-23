"""验证表源按键定向写入、原文件保留与漂移拒绝，不调用 Unity 或正式导出。"""

import contextlib
import io
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

TOOLS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TOOLS))
import nova_table_source as table
import openpyxl


class TableSourceTests(unittest.TestCase):
    def test_huge_numeric_request_is_structured_block_and_secret_aliases_are_hidden(self):
        """超大 JSON 整数不能使浮点转换溢出退出；SDK 常见密钥别名不进入查询或计划值。"""
        for kind in ("float", "double"):
            with self.subTest(kind=kind):
                with self.assertRaises(table.TableSourceError):
                    table._type_value(10 ** 1000, kind)
        for field in ("AFDevKey", "m_SdkKey", "WebhookUrl", "AndroidKeystorePassword"):
            self.assertTrue(table._private(field), field)

    def setUp(self):
        """每例使用隔离消费者目录，测试结束自动清除工作簿与状态文件。"""
        self.temp = tempfile.TemporaryDirectory(prefix="nova-table-source-test-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.source = "Assets/Tables/TableList.xlsx"
        target = self.root / self.source
        target.parent.mkdir(parents=True)
        repository = TOOLS.parents[3]
        shutil.copyfile(repository / "Assets/Samples/MainDemo/Excels/Tables/TableList.xlsx", target)
        self.plan_path = "Library/Nova/AgentActions/TableSource/test-plan.json"
        self.receipt_path = "Library/Nova/AgentActions/TableSource/test-receipt.json"

    def request(self, source=None, **cell_overrides):
        """从当前实际源摘要创建测试请求，不依赖 Excel 行号定位。"""
        source = source or self.source
        return {"source": source, "sourceHash": table._hash((self.root / source).read_bytes()),
                "sourceRoots": ["Assets/Tables"], "generatedRoots": ["Assets/Output"],
                "cells": [{"sheet": "Demo_Item", "rowKey": "Id", "key": 1001,
                           "column": "Price", "type": "int", "beforeValue": 100,
                           "value": 120, **cell_overrides}]}

    def fields_request(self, fields=None):
        """隔离项目用唯一正式自描述映射创建字段追加请求，不修改主仓真实 Schema 或 Excel。"""
        schema = "Assets/Tables/tables.xml"
        (self.root / schema).write_text('<schema><table name="TbDemo_Item" value="Demo_Item" input="Demo_Item@TableList.xlsx" mode="list" readSchemaFromFile="true"/></schema>', encoding="utf-8")
        probe = table.inspect_source(self.root, self.source, layout="table-fields", schema=schema,
                                     table="TbDemo_Item", sheet="Demo_Item", row_key="Id", row_key_type="int")
        return {"source": self.source, "sourceHash": probe["sourceHash"], "sourceRoots": ["Assets/Tables"],
                "generatedRoots": ["Assets/Output"], "layout": "table-fields", "schema": schema,
                "schemaHash": probe["schemaHash"], "table": "TbDemo_Item", "sheet": "Demo_Item",
                "rowKey": "Id", "rowKeyType": "int", "addFields": fields or [
                    {"name": "ReviewInt", "type": "int", "comment": "整数验证", "default": 7},
                    {"name": "ReviewLong", "type": "long", "comment": "大整数验证", "default": 2**63 - 1},
                    {"name": "ReviewFloat", "type": "float", "comment": "浮点验证", "default": 1.25},
                    {"name": "ReviewDouble", "type": "double", "comment": "双精度验证", "default": 2.5},
                    {"name": "ReviewBool", "type": "bool", "comment": "布尔验证", "default": True},
                    {"name": "ReviewText", "type": "string", "comment": "文本验证", "default": "回归验证"}]}

    def test_table_fields_all_six_types_defaults_and_original_content_preservation(self):
        """六种普通类型同计划追加全部原行，64位精确值及旧样式/其他页公式/ZIP资源均保留。"""
        source = self.root / self.source
        book = openpyxl.load_workbook(source)
        book["Demo_Item"]["E5"].number_format = "0.00"
        book["ListA"]["G5"] = "=1+2"
        book.save(source); book.close()
        with zipfile.ZipFile(source, "a") as archive:
            archive.writestr("customXml/owned-note.xml", "<note>原字节</note>")
        request = self.fields_request()
        before = source.read_bytes()
        planned = self.planned(request)
        self.assertEqual(planned["payload"]["schemaVersion"], 1)
        self.assertEqual(len(planned["payload"]["cells"]), 78)
        self.assertEqual(self.execute(planned)["status"], "success")
        self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "success")
        with zipfile.ZipFile(io.BytesIO(before)) as old, zipfile.ZipFile(source) as new:
            self.assertEqual([name for name in old.namelist() if old.read(name) != new.read(name)], ["xl/worksheets/sheet1.xml"])
            old_records = table._zip_records(before, old)[0]
            new_records = table._zip_records(source.read_bytes(), new)[0]
            self.assertTrue(all(old_records[name] == new_records[name] for name in old_records if name != "xl/worksheets/sheet1.xml"))
        book = openpyxl.load_workbook(source)
        sheet = book["Demo_Item"]
        self.assertEqual([sheet.cell(2, col).value for col in range(6, 12)], [field["name"] for field in request["addFields"]])
        for row in range(5, 15):
            values = [sheet.cell(row, col).value for col in range(6, 12)]
            self.assertEqual(values, [7, str(2**63 - 1), 1.25, 2.5, True, "回归验证"])
        self.assertEqual(sheet["E5"].value, 100)
        self.assertEqual(sheet["E5"].number_format, "0.00")
        self.assertEqual(book["ListA"]["G5"].value, "=1+2")
        book.close()
        self.assertEqual(table._table_fields_preservation(table._load(self.root, self.source), planned["payload"]["cells"]), planned["payload"]["preservationProof"])
        with self.assertRaises(table.TableSourceError):
            self.execute(planned)
        self.assertFalse((self.root / "Library/Nova/AgentActions/TableSource/write-lock.json").exists())
        self.assertEqual(list(source.parent.glob(".nova-table-source-*")), [])

    def test_table_fields_explicit_per_key_values_and_frozen_default_can_mix(self):
        """每个字段选择全集行键或默认值，逐行绑定真实键而不是列表顺序或Excel行号。"""
        request = self.fields_request()
        keys = table.inspect_source(self.root, self.source, layout="table-fields", schema=request["schema"],
                                    table=request["table"], sheet="Demo_Item", row_key="Id", row_key_type="int")["sheets"][0]["keys"]
        request["addFields"] = [{"name": "Quality", "type": "int", "comment": "逐行验证", "values": [{"key": item["key"], "value": i} for i, item in reversed(list(enumerate(keys)))]},
                                {"name": "Visible", "type": "bool", "comment": "默认验证", "default": False}]
        self.assertEqual(self.execute(self.planned(request))["status"], "success")
        book = openpyxl.load_workbook(self.root / self.source)
        self.assertEqual([book["Demo_Item"].cell(item["row"], 6).value for item in keys], list(range(10)))
        self.assertTrue(all(book["Demo_Item"].cell(item["row"], 7).value is False for item in keys)); book.close()
        self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "success")

    def test_table_fields_reject_bad_types_empty_defaults_and_ambiguous_specs(self):
        """类型、普通名称、注释和值必须明示，空值/复合Schema/凭据/公式与混合模式写前阻断。"""
        request = self.fields_request()
        base = {"name": "Quality", "type": "int", "comment": "品质", "default": 1}
        invalid = [{**base, "name": "Id"}, {**base, "name": "ClientToken"}, {**base, "name": "Name[]"},
                   {**base, "type": "enum"}, {**base, "type": []}, {**base, "default": "1"},
                   {**base, "default": True}, {**base, "default": None}, {**base, "default": 2**31},
                   {**base, "comment": ""}, {**base, "unknown": True}, {**base, "values": []},
                   {"name": "Quality", "type": "int", "comment": "品质"},
                   {**base, "type": "long", "default": 2**63}, {**base, "type": "float", "default": 1e39},
                   {**base, "type": "double", "default": float("inf")},
                   {**base, "type": "string", "default": ""}, {**base, "type": "string", "default": "=SUM(A1)"}]
        before = (self.root / self.source).read_bytes()
        for field in invalid:
            with self.subTest(field=field), self.assertRaises(table.TableSourceError):
                table.plan_patch(self.root, {**request, "addFields": [field]})
        with self.assertRaises(table.TableSourceError):
            table.plan_patch(self.root, {**request, "addFields": [base, base]})
        self.assertEqual((self.root / self.source).read_bytes(), before)

    def test_table_fields_reject_duplicate_missing_unknown_or_wrong_typed_keys(self):
        """逐行字段拒绝重复、缺失、未知和错类型行键，不能以default掩盖不完整values。"""
        request = self.fields_request()
        keys = table._table_fields_structure(table._load(self.root, self.source), "Demo_Item", "Id", "int")["keys"]
        values = [{"key": item["key"], "value": i} for i, item in enumerate(keys)]
        variants = [values[:-1], [values[0]] * len(values), [{"key": 999999, "value": 1}, *values[1:]],
                    [{"key": str(values[0]["key"]), "value": 1}, *values[1:]], [{"key": values[0]["key"], "value": "1"}, *values[1:]]]
        for supplied in variants:
            with self.subTest(supplied=supplied), self.assertRaises(table.TableSourceError):
                table.plan_patch(self.root, {**request, "addFields": [{"name": "Quality", "type": "int", "comment": "品质", "values": supplied}]})

    def test_table_fields_require_unique_self_described_schema_and_refuse_schema_drift(self):
        """正式Schema唯一映射/readSchemaFromFile及hash共同冻结，漂移不能消费计划或生成Receipt。"""
        request = self.fields_request()
        path = self.root / request["schema"]
        original = path.read_text()
        for altered in [original.replace('readSchemaFromFile="true"', 'readSchemaFromFile="false"'),
                        original.replace("Demo_Item@TableList.xlsx", "ListA@TableList.xlsx"),
                        original.replace("</schema>", original[8:-9] + "</schema>")]:
            with self.subTest(altered=altered), self.assertRaises(table.TableSourceError):
                path.write_text(altered, encoding="utf-8")
                table.inspect_source(self.root, self.source, layout="table-fields", schema=request["schema"], table="TbDemo_Item", sheet="Demo_Item", row_key="Id", row_key_type="int")
        path.write_text(original, encoding="utf-8")
        planned = self.planned(request)
        path.write_text(original + "\n", encoding="utf-8")
        with self.assertRaises(table.TableSourceError):
            self.execute(planned)
        self.assertFalse(table._read_state(self.root, self.plan_path)["executionStarted"])
        self.assertFalse((self.root / self.receipt_path).exists())

    def test_table_fields_refuse_formulas_merges_duplicate_markers_and_occupied_columns(self):
        """结构风险与既有相邻内容不能以追加操作覆盖，不能猜非自描述或复杂字段布局。"""
        source = self.root / self.source
        original = source.read_bytes()
        def formula(sheet): sheet["H14"] = "=1+2"
        def merge(sheet): sheet.merge_cells("G2:H2")
        def duplicate(sheet): sheet["A14"] = "##var"
        def occupied(sheet): sheet["F1"] = "已有非字段说明"
        def complex_type(sheet): sheet["E3"] = "array,int"
        def duplicate_key(sheet): sheet["B6"] = sheet["B5"].value
        def no_types(sheet): sheet["A3"] = ""
        for mutation in (formula, merge, duplicate, occupied, complex_type, duplicate_key, no_types):
            source.write_bytes(original)
            book = openpyxl.load_workbook(source); mutation(book["Demo_Item"]); book.save(source); book.close()
            with self.subTest(mutation=mutation.__name__), self.assertRaises(table.TableSourceError):
                request = self.fields_request()
                table.plan_patch(self.root, request)

    def test_table_fields_tampered_derived_write_set_and_verify_old_content_are_rejected(self):
        """重算planId也不能改变派生地址/样式；伪造整文件SHA不能隐藏旧数据被改。"""
        request = self.fields_request()
        planned = self.planned(request)
        tampered = json.loads(json.dumps(planned))
        tampered["payload"]["cells"][0]["styleFrom"] = "A1"
        tampered["planId"] = table._hash(table._canonical(tampered["payload"]))
        with self.assertRaises(table.TableSourceError):
            table._validate_plan(self.root, tampered)
        self.assertEqual(self.execute(planned)["status"], "success")
        source = self.root / self.source
        book = openpyxl.load_workbook(source); book["Demo_Item"]["E5"] = 999; book.save(source); book.close()
        receipt = table._read_state(self.root, self.receipt_path)
        receipt["expectedFileHash"] = table._hash(source.read_bytes())
        table._atomic_write(self.root, self.receipt_path, table._canonical(receipt), state=True)
        before = source.read_bytes()
        verified = table.verify_receipt(self.root, self.receipt_path)
        self.assertEqual(verified["status"], "partial")
        self.assertIn("table_original_content_changed", verified["issues"])
        self.assertEqual(source.read_bytes(), before)

    def test_table_fields_cli_inspect_and_failed_source_write_are_single_use(self):
        """新inspect参数使用同CLI；写入失败保留partial收据和原文件，不能重放或留下锁/tmp。"""
        request = self.fields_request()
        with contextlib.redirect_stdout(io.StringIO()) as output:
            code = table.main(["--project-root", str(self.root), "inspect", "--source", self.source,
                               "--layout", "table-fields", "--schema", request["schema"], "--table", request["table"],
                               "--sheet", "Demo_Item", "--row-key", "Id", "--row-key-type", "int"])
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(output.getvalue())["schemaHash"], request["schemaHash"])
        planned = self.planned(request)
        original = (self.root / self.source).read_bytes()
        writer = table._atomic_write
        def fail_source(root, relative, data, **kwargs):
            """仅模拟目标源写失败，计划/收据写仍复用正式原子入口。"""
            if relative == self.source: raise OSError("fixture-source-write")
            return writer(root, relative, data, **kwargs)
        with patch.object(table, "_atomic_write", side_effect=fail_source):
            self.assertEqual(self.execute(planned)["status"], "partial")
        self.assertEqual((self.root / self.source).read_bytes(), original)
        with self.assertRaises(table.TableSourceError): self.execute(planned)
        self.assertFalse((self.root / "Library/Nova/AgentActions/TableSource/write-lock.json").exists())
        self.assertEqual(list((self.root / self.source).parent.glob(".nova-table-source-*")), [])

    def test_table_fields_source_drift_blocks_and_schema_drift_after_execute_is_partial(self):
        """源在Plan后变化不得消费；正式写后Schema变化只读验收partial，不自动改源或重放。"""
        request = self.fields_request()
        planned = self.planned(request)
        source = self.root / self.source
        original = source.read_bytes()
        source.write_bytes(original + b"\n")
        with self.assertRaises(table.TableSourceError): self.execute(planned)
        self.assertFalse(table._read_state(self.root, self.plan_path)["executionStarted"])
        source.write_bytes(original)
        self.assertEqual(self.execute(planned)["status"], "success")
        schema = self.root / request["schema"]
        schema.write_bytes(schema.read_bytes() + b"\n")
        before = source.read_bytes()
        result = table.verify_receipt(self.root, self.receipt_path)
        self.assertEqual(result["status"], "partial")
        self.assertIn("table_schema_changed", result["issues"])
        self.assertEqual(source.read_bytes(), before)

    def planned(self, request=None):
        """建立并保存只读计划，随后测试正式一次性执行。"""
        result = table.plan_patch(self.root, request or self.request())
        table._atomic_write(self.root, self.plan_path, table._canonical(result), state=True)
        return result

    def execute(self, planned):
        """仅用返回计划 ID 执行固定源码编辑入口。"""
        return table.execute_plan(self.root, self.plan_path, planned["planId"], self.receipt_path)

    def test_state_cache_is_fixed_and_never_a_table_source(self):
        """仅固定缓存根放行状态 JSON，普通 Library 表源、其他缓存目录和状态链接仍被拒绝。"""
        table._atomic_write(self.root, self.plan_path, b"{}", state=True, replace=False)
        self.assertEqual(table._read_state(self.root, self.plan_path), {})
        self.assertFalse((self.root / ".agents").exists())
        for relative in (".agents/nova-table-source/plan.json", "Library/Other/plan.json",
                         "Library/Nova/AgentActions/TableSourceElse/plan.json",
                         table.STATE + "../plan.json", table.STATE + "plan.csv"):
            with self.subTest(relative=relative), self.assertRaises(table.TableSourceError):
                table._relative_path(self.root, relative, state=True, exists=False)
        for relative in ("Library/Tables/Source.xlsx", table.STATE + "Source.xlsx"):
            with self.subTest(relative=relative), self.assertRaises(table.TableSourceError):
                table._relative_path(self.root, relative, exists=False)
        (self.root / table.STATE / "linked.json").symlink_to(self.root / self.plan_path)
        with self.assertRaises(table.TableSourceError):
            table._read_state(self.root, table.STATE + "linked.json")

    def test_real_demo_key_patch_preserves_other_zip_parts_and_cells(self):
        """真实 Demo Id1001 Price100→120，其他 sheet/style/content 与邻行不变。"""
        source = self.root / self.source
        before = source.read_bytes()
        probe = table.inspect_source(self.root, self.source, sheet="Demo_Item", row_key="Id", key=1001, column="Price")
        self.assertEqual(probe["cell"]["beforeValue"], 100)
        planned = self.planned()
        self.assertEqual(self.execute(planned)["status"], "success")
        with zipfile.ZipFile(io.BytesIO(before)) as old, zipfile.ZipFile(source) as new:
            changed = [name for name in old.namelist() if old.read(name) != new.read(name)]
            self.assertEqual(changed, ["xl/worksheets/sheet1.xml"])
        book = openpyxl.load_workbook(source)
        self.assertEqual(book["Demo_Item"]["E5"].value, 120)
        self.assertEqual(book["Demo_Item"]["E6"].value, 200)
        book.close()
        saved = source.read_bytes()
        self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "success")
        self.assertEqual(source.read_bytes(), saved)
        with self.assertRaises(table.TableSourceError):
            self.execute(planned)
        self.assertFalse((self.root / "Library/Nova/AgentActions/TableSource/write-lock.json").exists())
        self.assertEqual(list(source.parent.glob(".nova-table-source-*")), [])

    def test_formula_styles_and_xml_extensions_are_preserved(self):
        """对值编辑不重新保存工作簿，邻格公式、样式及附加 ZIP 资源逐字保留。"""
        source = self.root / self.source
        book = openpyxl.load_workbook(source)
        book["Demo_Item"]["F5"] = "=E5*2"
        book["Demo_Item"]["E5"].number_format = "0.00"
        book.save(source)
        book.close()
        with zipfile.ZipFile(source, "a") as archive:
            archive.writestr("customXml/nova-note.xml", "<note>保留</note>")
        self.assertEqual(self.execute(self.planned())["status"], "success")
        book = openpyxl.load_workbook(source)
        self.assertEqual(book["Demo_Item"]["F5"].value, "=E5*2")
        self.assertEqual(book["Demo_Item"]["E5"].number_format, "0.00")
        book.close()
        with zipfile.ZipFile(source) as archive:
            self.assertEqual(archive.read("customXml/nova-note.xml"), "<note>保留</note>".encode())

    def test_wrong_before_hash_type_key_and_duplicates_are_rejected(self):
        """冻结原值/hash和类型；错误键与重复目标不能形成可执行计划。"""
        for overrides in [{"beforeValue": 99}, {"value": "120"}, {"value": True},
                          {"key": 999}, {"type": "string"}, {"column": "Id"}, {"value": 2**31}]:
            with self.subTest(overrides=overrides), self.assertRaises(table.TableSourceError):
                table.plan_patch(self.root, self.request(**overrides))
        request = self.request()
        request["sourceHash"] = "0" * 64
        with self.assertRaises(table.TableSourceError):
            table.plan_patch(self.root, request)
        request = self.request()
        request["cells"] *= 2
        with self.assertRaises(table.TableSourceError):
            table.plan_patch(self.root, request)
        book = openpyxl.load_workbook(self.root / self.source)
        book["Demo_Item"]["B6"] = 1001
        book.save(self.root / self.source)
        book.close()
        with self.assertRaises(table.TableSourceError):
            table.plan_patch(self.root, self.request())

    def test_source_drift_and_foreign_confirmation_leave_source_untouched(self):
        """计划后任何文件漂移或错误 token 都阻止写入，不能只比较目标原值。"""
        planned = self.planned()
        source = self.root / self.source
        with self.assertRaises(table.TableSourceError):
            table.execute_plan(self.root, self.plan_path, "foreign", self.receipt_path)
        with zipfile.ZipFile(source, "a") as archive:
            archive.writestr("customXml/changed.xml", "<changed/>")
        before = source.read_bytes()
        with self.assertRaises(table.TableSourceError):
            self.execute(planned)
        self.assertEqual(source.read_bytes(), before)
        self.assertFalse((self.root / self.receipt_path).exists())

    def test_path_generated_vendor_symlink_and_sensitive_fields_are_rejected(self):
        """拒绝逃逸、链接、厂商和生成物，不回显身份/凭据值。"""
        for source in ["../outside.xlsx", "/tmp/outside.xlsx", "Assets/Vendor/Table.xlsx", "Packages/SDK/Table.xlsx"]:
            with self.subTest(source=source), self.assertRaises(table.TableSourceError):
                table.inspect_source(self.root, source)
        request = self.request()
        request["generatedRoots"] = ["Assets/Tables"]
        with self.assertRaises(table.TableSourceError):
            table.plan_patch(self.root, request)
        (self.root / "Assets/Tables/linked.xlsx").symlink_to(self.root / self.source)
        with self.assertRaises(table.TableSourceError):
            table.inspect_source(self.root, "Assets/Tables/linked.xlsx")
        book = openpyxl.load_workbook(self.root / self.source)
        book["Demo_Item"]["C2"] = "ClientToken"
        book.save(self.root / self.source)
        book.close()
        with self.assertRaises(table.TableSourceError):
            table.inspect_source(self.root, self.source, sheet="Demo_Item", row_key="Id", key=1001, column="ClientToken")

    def test_csv_bom_multiline_quotes_and_blank_records_preserve_exact_other_bytes(self):
        """CSV 只替换目标字段，BOM/换行/转义/空行和相邻数据原样保留。"""
        source = "Assets/Tables/Items.csv"
        original = '\ufeff\r\n##var,Id,Name,Price\r\n##type,int,string,int\r\n,1001,"demo,\r\n""sword""",100\r\n\r\n,1002,shield,200\r\n'
        (self.root / source).write_bytes(original.encode())
        request = self.request(source)
        request["cells"][0]["sheet"] = "csv"
        self.assertEqual(self.execute(self.planned(request))["status"], "success")
        self.assertEqual((self.root / source).read_bytes(), original.replace('",100\r\n', '",120\r\n').encode())
        self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "success")

    def test_verify_changed_source_is_partial_and_never_writes(self):
        """Verify 发现源或目标漂移时只报告 partial，不回滚和重新执行。"""
        self.execute(self.planned())
        source = self.root / self.source
        book = openpyxl.load_workbook(source)
        book["Demo_Item"]["E5"] = 130
        book.save(source)
        book.close()
        before = source.read_bytes()
        result = table.verify_receipt(self.root, self.receipt_path)
        self.assertEqual(result["status"], "partial")
        self.assertIn("cell_value_or_key_changed", result["issues"])
        self.assertEqual(source.read_bytes(), before)

    def test_macro_and_missing_reader_block_without_output_leak(self):
        """当前不支持宏；缺 openpyxl 时准确 blocked，不自动安装或修改源。"""
        source = self.root / self.source
        with zipfile.ZipFile(source, "a") as archive:
            archive.writestr("xl/vbaProject.bin", b"macro")
        with self.assertRaises(table.TableSourceError):
            table.plan_patch(self.root, self.request())
        shutil.copyfile(TOOLS.parents[3] / "Assets/Samples/MainDemo/Excels/Tables/TableList.xlsx", source)
        with patch.dict(sys.modules, {"openpyxl": None}), self.assertRaisesRegex(table.TableSourceError, "openpyxl"):
            table.inspect_source(self.root, self.source)

    def test_cli_outputs_summary_and_excludes_data_values(self):
        """CLI 可读写固定缓存请求、计划与回执，且不回显 before/value。"""
        request_relative = table.STATE + "request.json"
        request_path = self.root / request_relative
        request_path.parent.mkdir(parents=True)
        request_path.write_text(json.dumps(self.request()))
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            code = table.main(["--project-root", str(self.root), "plan", "--request", request_relative,
                               "--plan", self.plan_path])
        self.assertEqual(code, 0)
        planned = json.loads(output.getvalue())
        self.assertEqual(planned["plan"], self.plan_path)
        self.assertNotIn("beforeValue", planned["cells"][0])
        self.assertNotIn("value", planned["cells"][0])
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(table.main(["--project-root", str(self.root), "execute", "--plan", self.plan_path,
                                        "--confirmation-token", planned["planId"], "--receipt", self.receipt_path]), 0)
            self.assertEqual(table.main(["--project-root", str(self.root), "verify", "--receipt", self.receipt_path]), 0)
        self.assertFalse((self.root / ".agents").exists())

    def test_plain_csv_requires_types_and_rejects_formula_strings(self):
        """普通 CSV 不推断 ID 数字类型，公式样式字符串不能由编辑新增。"""
        source = "Assets/Tables/Plain.csv"
        (self.root / source).write_text("Id;Name;Price\n1001;sword;100\n1002;shield;200\n")
        request = self.request(source)
        request["cells"][0]["sheet"] = "csv"
        with self.assertRaises(table.TableSourceError):
            table.plan_patch(self.root, request)
        request["cells"][0]["rowKeyType"] = "int"
        self.assertEqual(self.execute(self.planned(request))["status"], "success")
        self.assertEqual((self.root / source).read_text(), "Id;Name;Price\n1001;sword;120\n1002;shield;200\n")
        for value in ["=1+2", " +SUM(A1)", "@formula", "-formula"]:
            edit = self.request(source)
            edit["cells"][0].update(sheet="csv", rowKeyType="int", column="Name", type="string", beforeValue="sword", value=value)
            with self.subTest(value=value), self.assertRaises(table.TableSourceError):
                table.plan_patch(self.root, edit)

    def test_xlsx_literal_string_bool_long_and_blank_cell(self):
        """多字段写入保留文本公式、bool 和 64 位整数精度，并允许已有空格。"""
        source = self.root / self.source
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Items"
        sheet.append(["##var", "Id", "Title", "Enabled", "Amount", "Note"])
        sheet.append(["##type", "int", "string", "bool", "long", "string"])
        sheet.append([None, 1, "before", False, 0, None])
        sheet["F3"].number_format = "@"
        book.save(source)
        book.close()
        request = self.request()
        request["cells"] = [
            {"sheet": "Items", "rowKey": "Id", "key": 1, "column": column, "type": kind, "beforeValue": before, "value": after}
            for column, kind, before, after in [("Title", "string", "before", "=1+2"), ("Enabled", "bool", False, True),
                                               ("Amount", "long", 0, 2**63 - 1), ("Note", "string", None, "文本")]
        ]
        self.assertEqual(self.execute(self.planned(request))["status"], "success")
        book = openpyxl.load_workbook(source)
        self.assertEqual(book["Items"]["C3"].value, "=1+2")
        self.assertEqual(book["Items"]["C3"].data_type, "s")
        self.assertIs(book["Items"]["D3"].value, True)
        self.assertEqual(book["Items"]["E3"].value, str(2**63 - 1))
        self.assertEqual(book["Items"]["F3"].number_format, "@")
        book.close()
        self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "success")

    def test_formula_target_and_ancestor_symlink_are_rejected(self):
        """已有公式不能覆盖，项目根也不能经祖先目录链接绕过边界。"""
        source = self.root / self.source
        book = openpyxl.load_workbook(source)
        book["Demo_Item"]["E5"] = "=50*2"
        book.save(source)
        book.close()
        with self.assertRaises(table.TableSourceError):
            table.plan_patch(self.root, self.request())
        (self.root / "nested").mkdir()
        (self.root / "linked-root").symlink_to(self.root, target_is_directory=True)
        with self.assertRaises(table.TableSourceError):
            table._project_root(self.root / "linked-root/nested")

    def test_atomic_write_failure_keeps_receipt_and_consumes_plan_without_replay(self):
        """源替换失败后留下恢复证据与已消费计划，Verify 只读并报告 partial。"""
        planned = self.planned()
        source = self.root / self.source
        before = source.read_bytes()
        original_replace = table.os.replace
        def fail_source_replace(src, target):
            """仅模拟真实源的原子替换失败，状态写入仍可完成。"""
            if Path(target) == source:
                raise OSError("测试源替换失败")
            return original_replace(src, target)
        with patch.object(table.os, "replace", side_effect=fail_source_replace):
            result = self.execute(planned)
        self.assertEqual(result["status"], "partial")
        self.assertEqual(source.read_bytes(), before)
        self.assertTrue(json.loads((self.root / self.plan_path).read_text())["executionStarted"])
        self.assertFalse(json.loads((self.root / self.receipt_path).read_text())["completed"])
        verified = table.verify_receipt(self.root, self.receipt_path)
        self.assertEqual(verified["status"], "partial")
        self.assertIn("completion_evidence_missing", verified["issues"])
        with self.assertRaises(table.TableSourceError):
            self.execute(planned)
        self.assertFalse((self.root / "Library/Nova/AgentActions/TableSource/write-lock.json").exists())
        self.assertEqual(list(source.parent.glob(".nova-table-source-*")), [])

    def test_invalid_receipt_and_plan_structure_are_rejected(self):
        """Verify 不接受空写集、伪完成标记或缺少摘要的手改状态。"""
        self.execute(self.planned())
        receipt = json.loads((self.root / self.receipt_path).read_text())
        for change in ["empty_cells", "unstarted", "missing_hash"]:
            changed = json.loads(json.dumps(receipt))
            if change == "empty_cells":
                changed["plan"]["payload"]["cells"] = []
                changed["plan"]["planId"] = table._hash(table._canonical(changed["plan"]["payload"]))
            elif change == "unstarted":
                changed["plan"]["executionStarted"] = False
            else:
                changed["expectedFileHash"] = ""
            (self.root / self.receipt_path).write_bytes(table._canonical(changed))
            with self.subTest(change=change), self.assertRaises(table.TableSourceError):
                table.verify_receipt(self.root, self.receipt_path)

    def tracks_request(self, *, column="属性说明", property_name="uwr_path", value="只记录不含 query 的路径。"):
        """隔离复制真实 Framework Tracks 源，构造事件/属性复合身份请求。"""
        source = "Assets/Tables/Tracks.xlsx"
        target = self.root / source
        if not target.exists():
            shutil.copyfile(TOOLS.parents[3] / "Assets/Framework/Tracks/Tracks.xlsx", target)
        probe = table.inspect_source(self.root, source, layout="tracks-events", sheet="打点表-Network",
                                     event_name="uwr_request_start", property_name=property_name, column=column)
        return {"source": source, "sourceHash": probe["sourceHash"], "sourceRoots": ["Assets/Tables"],
                "generatedRoots": ["Assets/Output"], "layout": "tracks-events",
                "cells": [{k: probe["cell"][k] for k in ("sheet", "eventName", "propertyName", "column", "type", "beforeValue")} | {"value": value}]}

    def test_real_tracks_leaf_preserves_identity_targets_styles_and_other_zip_parts(self):
        """真实 Network 属性说明定向修改，所有其它 ZIP 部件、目标和样式不变。"""
        request = self.tracks_request()
        source = self.root / request["source"]
        before = source.read_bytes()
        old_book = openpyxl.load_workbook(io.BytesIO(before))
        old_sheet = old_book["打点表-Network"]
        old_styles = [(cell.coordinate, cell.style_id) for row in old_sheet for cell in row]
        old_merges = sorted(map(str, old_sheet.merged_cells.ranges))
        old_targets = [cell.value for cell in old_sheet["H"]]
        old_book.close()
        planned = self.planned(request)
        self.assertEqual(planned["payload"]["cells"][0]["address"], "F10")
        self.assertEqual(self.execute(planned)["status"], "success")
        with zipfile.ZipFile(io.BytesIO(before)) as old, zipfile.ZipFile(source) as new:
            self.assertEqual([name for name in old.namelist() if old.read(name) != new.read(name)], ["xl/worksheets/sheet5.xml"])
        book = openpyxl.load_workbook(source)
        sheet = book["打点表-Network"]
        self.assertEqual(sheet["F10"].value, request["cells"][0]["value"])
        self.assertEqual(sheet["D10"].value, "uwr_path")
        self.assertEqual([cell.value for cell in sheet["H"]], old_targets)
        self.assertEqual([(cell.coordinate, cell.style_id) for row in sheet for cell in row], old_styles)
        self.assertEqual(sorted(map(str, sheet.merged_cells.ranges)), old_merges)
        book.close()
        self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "success")

    def test_tracks_merged_description_targets_anchor_and_duplicate_anchor_blocks(self):
        """从合并区任一属性定位 C 首格，多属性指向同首格时拒绝重复写入。"""
        request = self.tracks_request(column="事件说明", value="一条逻辑请求开始，仅登记一次。")
        planned = self.planned(request)
        cell = planned["payload"]["cells"][0]
        self.assertEqual((cell["address"], cell["selectedPropertyRow"], cell["eventStartRow"], cell["mergedRange"]), ("C2", 10, 2, "C2:C20"))
        self.assertEqual(self.execute(planned)["status"], "success")
        book = openpyxl.load_workbook(self.root / request["source"])
        self.assertEqual(book["打点表-Network"]["C2"].value, request["cells"][0]["value"])
        self.assertIsNone(book["打点表-Network"]["C10"].value)
        book.close()
        self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "success")
        duplicate = self.tracks_request(column="事件说明")
        duplicate["cells"].append({**duplicate["cells"][0], "propertyName": "uwr_host"})
        with self.assertRaises(table.TableSourceError):
            table.plan_patch(self.root, duplicate)

    def test_tracks_duplicate_property_event_and_invalid_merge_block(self):
        """复合身份重复或跨列合并不能形成计划，不按第一条记录猜测。"""
        original = self.tracks_request()
        source = self.root / original["source"]
        before = source.read_bytes()
        for kind in ("property", "event", "merge"):
            source.write_bytes(before)
            book = openpyxl.load_workbook(source)
            sheet = book["打点表-Network"]
            if kind == "property":
                sheet["D9"] = "uwr_path"
            elif kind == "event":
                sheet["A21"] = "uwr_request_start"
            else:
                sheet.merge_cells("G2:G3")
            book.save(source)
            book.close()
            request = {**original, "sourceHash": table._hash(source.read_bytes())}
            with self.subTest(kind=kind), self.assertRaises(table.TableSourceError):
                table.plan_patch(self.root, request)

    def test_tracks_identity_target_sensitive_and_unknown_layout_are_rejected(self):
        """身份与目标保持只读，敏感属性和值不公开，未知布局不回退普通表。"""
        original = self.tracks_request()
        for changes in ({"column": "目标"}, {"column": "事件名"}, {"column": "属性名"},
                        {"propertyName": "device_id"}, {"value": "Bearer never-output"}, {"type": "int"},
                        {"propertyName": "无固定属性"}):
            request = {**original, "cells": [{**original["cells"][0], **changes}]}
            with self.subTest(changes=changes), self.assertRaises(table.TableSourceError):
                table.plan_patch(self.root, request)
        with self.assertRaises(table.TableSourceError):
            table.plan_patch(self.root, {**original, "layout": "tracks"})
        with self.assertRaises(table.TableSourceError):
            table.inspect_source(self.root, original["source"], layout="tracks-events", sheet="默认用户属性说明")
        with self.assertRaises(table.TableSourceError):
            table.inspect_source(self.root, original["source"], layout="tracks-events", sheet="打点表-Network",
                                 event_name="uwr_request_start", property_name="access_token", column="属性说明")

    def test_tracks_source_and_merge_drift_reject_execute_and_verify_reports_partial(self):
        """执行前整源漂移阻断；执行后合并漂移只读报告 partial，不能回滚重放。"""
        request = self.tracks_request()
        planned = self.planned(request)
        source = self.root / request["source"]
        before = source.read_bytes()
        with zipfile.ZipFile(source, "a") as archive:
            archive.writestr("customXml/changed.xml", "<changed/>")
        with self.assertRaises(table.TableSourceError):
            self.execute(planned)
        source.write_bytes(before)
        self.assertEqual(self.execute(planned)["status"], "success")
        book = openpyxl.load_workbook(source)
        book["打点表-Network"].unmerge_cells("C2:C20")
        book.save(source)
        book.close()
        frozen = source.read_bytes()
        result = table.verify_receipt(self.root, self.receipt_path)
        self.assertEqual(result["status"], "partial")
        self.assertIn("cell_position_changed", result["issues"])
        self.assertEqual(source.read_bytes(), frozen)

    def localization_request(self, *, add_language=True, add_key=True):
        """隔离复制真实 Nova 多语言表，新增列与 Key 的所有译文均显式指定。"""
        source = "Assets/Tables/LocalizationTexts.xlsx"
        target = self.root / source
        if not target.exists():
            shutil.copyfile(TOOLS.parents[3] / "Assets/Samples/MainDemo/Excels/Localizations/Texts/LocalizationTexts.xlsx", target)
        probe = table.inspect_source(self.root, source, layout="localization-texts")
        sheet = probe["sheets"][0]
        request = {"source": source, "sourceHash": probe["sourceHash"], "sourceRoots": ["Assets/Tables"],
                   "generatedRoots": ["Assets/Output"], "layout": "localization-texts"}
        if add_language:
            request["addLanguages"] = [{"sheet": sheet["name"], "language": "Japanese",
                                        "values": [{"key": key, "value": "" if i == 0 else "訳文 " + str(i)}
                                                   for i, key in enumerate(sheet["keys"])]}]
        if add_key:
            translations = {"ChineseSimplified": "新增提示", "English": "New text"}
            if add_language:
                translations["Japanese"] = "新しい文"
            request["addKeys"] = [{"sheet": sheet["name"], "key": "Review_NewText",
                                   "description": "隔离验证", "translations": translations}]
        return request

    def test_localization_real_source_add_language_and_key_preserves_all_old_content(self):
        """真实表同计划新增日语和 Key，保留旧字节部件、样式、合并及旧单元格值。"""
        request = self.localization_request()
        source = self.root / request["source"]
        before = source.read_bytes()
        old = openpyxl.load_workbook(io.BytesIO(before))
        old_sheet = old["Demo_Localization"]
        old_values = {(cell.coordinate): (cell.value, cell.style_id) for row in old_sheet for cell in row
                      if cell.column <= 5 or cell.column >= 7}
        old_merges = list(map(str, old_sheet.merged_cells.ranges))
        old.close()
        planned = self.planned(request)
        self.assertEqual(planned["payload"]["layout"], "localization-texts")
        self.assertEqual(self.execute(planned)["status"], "success")
        with zipfile.ZipFile(io.BytesIO(before)) as old_zip, zipfile.ZipFile(source) as new_zip:
            self.assertEqual([name for name in old_zip.namelist() if old_zip.read(name) != new_zip.read(name)],
                             ["xl/worksheets/sheet1.xml"])
            old_cells = {m.group(0) for m in table.CELL_XML.finditer(old_zip.read("xl/worksheets/sheet1.xml").decode())
                         if not table.re.search(r'\br="F[0-9]+"', m.group("attrs"))}
            new_xml = new_zip.read("xl/worksheets/sheet1.xml").decode()
            self.assertTrue(all(cell in new_xml for cell in old_cells))
        book = openpyxl.load_workbook(source)
        sheet = book["Demo_Localization"]
        self.assertEqual({coordinate: (sheet[coordinate].value, sheet[coordinate].style_id) for coordinate in old_values}, old_values)
        self.assertEqual(list(map(str, sheet.merged_cells.ranges)), old_merges)
        self.assertEqual(sheet["F2"].value, "Japanese")
        self.assertEqual(sheet["F3"].value, "string")
        self.assertEqual(sheet["F5"].value, "")
        self.assertEqual(sheet["B68"].value, "Review_NewText")
        self.assertEqual(sheet["F68"].value, "新しい文")
        self.assertEqual(sheet["B68"].style_id, sheet["B9"].style_id)
        book.close()
        self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "success")
        with self.assertRaises(table.TableSourceError):
            self.execute(planned)
        self.assertFalse((self.root / "Library/Nova/AgentActions/TableSource/write-lock.json").exists())

    def test_localization_add_key_only_and_add_language_only_use_same_transaction(self):
        """两个独立目标都复用同一 Plan/Execute/Verify，不要求额外无关结构修改。"""
        for add_language, add_key in ((False, True), (True, False)):
            with self.subTest(add_language=add_language):
                request = self.localization_request(add_language=add_language, add_key=add_key)
                self.assertEqual(self.execute(self.planned(request))["status"], "success")
                self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "success")
                for relative in (self.plan_path, self.receipt_path, request["source"]):
                    (self.root / relative).unlink()

    def test_localization_structure_rejects_unknown_duplicate_missing_and_formula_inputs(self):
        """不能猜语言拼写、填充值或 Key，且结构计划不能夹带 scalar/insert 操作。"""
        original = self.localization_request()
        invalid = []
        for language in ("japanese", "NotALanguage", "Unspecified", "English"):
            request = json.loads(json.dumps(original))
            request["addLanguages"][0]["language"] = language
            invalid.append(request)
        request = json.loads(json.dumps(original))
        request["addLanguages"][0]["values"].pop()
        invalid.append(request)
        request = json.loads(json.dumps(original))
        request["addLanguages"][0]["values"][0]["key"] = request["addLanguages"][0]["values"][1]["key"]
        invalid.append(request)
        request = json.loads(json.dumps(original))
        request["addKeys"][0]["translations"].pop("Japanese")
        invalid.append(request)
        request = json.loads(json.dumps(original))
        request["addKeys"][0]["key"] = "Demo_Greet"
        invalid.append(request)
        request = json.loads(json.dumps(original))
        request["addKeys"][0]["translations"]["English"] = "=SUM(1,2)"
        invalid.append(request)
        request = json.loads(json.dumps(original))
        request["cells"] = []
        invalid.append(request)
        for request in invalid:
            with self.subTest(request_index=invalid.index(request)), self.assertRaises(table.TableSourceError):
                table.plan_patch(self.root, request)
        self.assertEqual(table._hash((self.root / original["source"]).read_bytes()), original["sourceHash"])

    def test_localization_occupied_column_merge_formula_and_bad_header_block_without_write(self):
        """新增语言只填相邻空白列，备注/合并/公式/未知表头不会被搬动或覆盖。"""
        original = self.localization_request()
        source = self.root / original["source"]
        before = source.read_bytes()
        for kind in ("occupied", "merge", "formula", "header", "duplicate_header"):
            source.write_bytes(before)
            book = openpyxl.load_workbook(source)
            sheet = book["Demo_Localization"]
            if kind == "occupied":
                sheet["F10"] = "保持备注"
            elif kind == "merge":
                sheet.merge_cells("F2:G2")
            elif kind == "formula":
                sheet["E5"] = "=1+1"
            elif kind == "header":
                sheet["G2"] = "UnknownLocale"
                sheet["G3"] = "string"
            else:
                sheet["F2"] = "#English"
                sheet["F3"] = "string"
            book.save(source)
            book.close()
            request = {**original, "sourceHash": table._hash(source.read_bytes())}
            with self.subTest(kind=kind), self.assertRaises(table.TableSourceError):
                table.plan_patch(self.root, request)
            self.assertEqual(table._hash(source.read_bytes()), request["sourceHash"])

    def test_localization_all_effective_sheets_require_same_added_languages(self):
        """语言新增覆盖工作簿全部有效 Sheet，避免正式专用导出才发现集合不一致。"""
        request = self.localization_request(add_key=False)
        source = self.root / request["source"]
        book = openpyxl.load_workbook(source)
        book.copy_worksheet(book["Demo_Localization"]).title = "Second"
        book.save(source)
        book.close()
        request["sourceHash"] = table._hash(source.read_bytes())
        with self.assertRaises(table.TableSourceError):
            table.plan_patch(self.root, request)
        request["addLanguages"].append({**request["addLanguages"][0], "sheet": "Second"})
        self.assertEqual(self.execute(self.planned(request))["status"], "success")
        self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "success")

    def test_localization_source_and_enum_drift_are_blocked_and_verify_never_repairs(self):
        """源或真实语言契约漂移阻止执行；执行后破坏 Key 的 Verify 只读报告 partial。"""
        request = self.localization_request()
        planned = self.planned(request)
        source = self.root / request["source"]
        before = source.read_bytes()
        with patch.object(table, "_localization_languages", return_value=({"Japanese"}, "0" * 64)):
            with self.assertRaises(table.TableSourceError):
                self.execute(planned)
        with zipfile.ZipFile(source, "a") as archive:
            archive.writestr("customXml/changed.xml", "<changed/>")
        with self.assertRaises(table.TableSourceError):
            self.execute(planned)
        source.write_bytes(before)
        self.assertEqual(self.execute(planned)["status"], "success")
        book = openpyxl.load_workbook(source)
        book["Demo_Localization"]["B68"] = "Renamed"
        book.save(source)
        book.close()
        frozen = source.read_bytes()
        self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "partial")
        self.assertEqual(source.read_bytes(), frozen)

    def test_localization_multiple_languages_keys_and_legacy_hash_headers(self):
        """兼容单个 # 语言表头，同时新增两语两 Key，旧列和新值顺序明确。"""
        request = self.localization_request()
        source = self.root / request["source"]
        book = openpyxl.load_workbook(source)
        book["Demo_Localization"]["E2"] = "#English"
        book.save(source)
        book.close()
        request["sourceHash"] = table._hash(source.read_bytes())
        request["addLanguages"].append({**request["addLanguages"][0], "language": "French",
                                        "values": [{**item, "value": "fr"} for item in request["addLanguages"][0]["values"]]})
        request["addKeys"][0]["translations"]["French"] = "Nouveau"
        request["addKeys"].append({**request["addKeys"][0], "key": "Review_NewText2"})
        self.assertEqual(self.execute(self.planned(request))["status"], "success")
        book = openpyxl.load_workbook(source)
        sheet = book["Demo_Localization"]
        self.assertEqual((sheet["E2"].value, sheet["F2"].value, sheet["G2"].value), ("#English", "Japanese", "French"))
        self.assertEqual((sheet["B68"].value, sheet["B69"].value, sheet["G69"].value),
                         ("Review_NewText", "Review_NewText2", "Nouveau"))
        book.close()
        self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "success")

    def test_localization_tampered_plan_limits_and_generated_boundary_are_rejected(self):
        """重算摘要也不能注入无效结构地址或额外写集，源码生成边界继续生效。"""
        request = self.localization_request()
        forbidden = {**request, "generatedRoots": ["Assets/Tables"]}
        with self.assertRaises(table.TableSourceError):
            table.plan_patch(self.root, forbidden)
        planned = self.planned(request)
        for change in ("row_limit", "empty", "extra_field", "unknown_style"):
            tampered = json.loads(json.dumps(planned))
            if change == "row_limit":
                tampered["payload"]["cells"][0]["row"] = 1048577
                tampered["payload"]["cells"][0]["address"] = "F1048577"
            elif change == "empty":
                tampered["payload"]["cells"] = []
            elif change == "extra_field":
                tampered["payload"]["cells"][0]["insertColumn"] = True
            else:
                tampered["payload"]["cells"][0]["styleFrom"] = "A1"
            tampered["planId"] = table._hash(table._canonical(tampered["payload"]))
            table._atomic_write(self.root, self.plan_path, table._canonical(tampered), state=True)
            with self.subTest(change=change), self.assertRaises(table.TableSourceError):
                self.execute(tampered)
        self.assertEqual(table._hash((self.root / request["source"]).read_bytes()), request["sourceHash"])
        self.assertFalse((self.root / self.receipt_path).exists())

    def test_localization_cli_summary_and_write_failure_consume_once(self):
        """结构 CLI 返回数据脱敏摘要；原子替换失败留 Receipt 且不得重放。"""
        request = self.localization_request()
        request_path = self.root / "Assets/Tables/localization-request.json"
        request_path.write_bytes(table._canonical(request))
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            self.assertEqual(table.main(["--project-root", str(self.root), "plan", "--request",
                                         "Assets/Tables/localization-request.json", "--plan", self.plan_path]), 0)
        planned_summary = json.loads(output.getvalue())
        self.assertNotIn("译文", output.getvalue())
        self.assertTrue(all(set(cell) == {"sheet", "type", "address"} for cell in planned_summary["cells"]))
        source = self.root / request["source"]
        before = source.read_bytes()
        original_replace = table.os.replace
        def fail_source_replace(src, target):
            """只模拟目标表源替换失败，不阻断写前状态凭据。"""
            if Path(target) == source:
                raise OSError("隔离替换失败")
            return original_replace(src, target)
        with patch.object(table.os, "replace", side_effect=fail_source_replace):
            result = table.execute_plan(self.root, self.plan_path, planned_summary["planId"], self.receipt_path)
        self.assertEqual(result["status"], "partial")
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(table.verify_receipt(self.root, self.receipt_path)["status"], "partial")
        with self.assertRaises(table.TableSourceError):
            table.execute_plan(self.root, self.plan_path, planned_summary["planId"], self.receipt_path)
        self.assertFalse((self.root / "Library/Nova/AgentActions/TableSource/write-lock.json").exists())
        self.assertEqual(list(source.parent.glob(".nova-table-source-*")), [])


if __name__ == "__main__":
    unittest.main()

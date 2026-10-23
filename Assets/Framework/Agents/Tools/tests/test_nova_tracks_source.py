"""Tracks 结构编辑的权威来源、整表事务和恢复边界回归。"""
import copy
import contextlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from unittest import mock

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import nova_tracks_source as tracks
import nova_table_source as table
import openpyxl


class TracksSourceTests(unittest.TestCase):
    """每个用例创建完整隔离登记项目，临时产物由 TemporaryDirectory 精确清理。"""
    def setUp(self):
        """创建实际九列事件源、说明页与源码证据，并冻结与 Unity 相同的来源快照。"""
        self.temp=tempfile.TemporaryDirectory(prefix='nova-tracks-tests-',dir=str(Path(tempfile.gettempdir()).resolve()))
        self.root=Path(self.temp.name)
        self.path=self.root/'Assets/Framework/Tracks/Tracks.xlsx'
        self.path.parent.mkdir(parents=True)
        book=openpyxl.Workbook();sheet=book.active;sheet.title='打点表-Network'
        sheet.append(list(tracks.HEADERS))
        sheet.append(['$MPRegister','注册','注册事件','level_id','int','当前关卡','注册成功','gravity','原备注'])
        sheet.append(['','','','scene','string','当前页面','触发时','tga','次备注'])
        sheet.merge_cells('A2:A3');sheet.merge_cells('B2:B3');sheet.merge_cells('C2:C3')
        second=book.create_sheet('说明页');second.append(['固定属性','说明']);second.append(['真实字段','中文说明'])
        book.save(self.path);book.close()
        # 未修改 ZIP 部件保持逐字节原样，不以 openpyxl.save 作为生产写入器。
        with zipfile.ZipFile(self.path,'a') as archive:
            archive.writestr('custom/keep.xml',b'<keep exact="yes"/>')
        self.code=self.root/'Assets/Game/Tracking.cs';self.code.parent.mkdir(parents=True)
        self.code.write_text('// 实际业务调用\nclass Tracking {}\n')
        self.registry=self.snapshot()
        self.request={'source':'Assets/Framework/Tracks/Tracks.xlsx','sourceHash':table._hash(self.path.read_bytes()),
                      'ownerRoot':'Assets/Framework','registry':self.registry,
                      'codeEvidence':[{'path':'Assets/Game/Tracking.cs','sha256':table._hash(self.code.read_bytes())}],
                      'operations':[{'kind':'normalizeStyle','sheet':'打点表-Network'}]}

    def tearDown(self):
        """销毁隔离项目与计划、Receipt、备份和源，避免测试留下状态或生成噪声。"""
        self.temp.cleanup()

    def snapshot(self):
        """测试根没有 Unity 解析包时按固定规则构造全部真实登记源与元数据。"""
        sources=[self.path]
        packages=sorted((self.root/'UPMPackages').glob('*')) if (self.root/'UPMPackages').exists() else []
        sources += [p/'Nova/Tracks/Tracks.xlsx' for p in packages if (p/'Nova/Tracks/Tracks.xlsx').is_file()]
        metadata=[self.root/'Packages/manifest.json',self.root/'Packages/packages-lock.json']
        metadata += [p.with_name(p.name+'.meta') for p in sources]
        metadata += [p/'package.json' for p in packages]
        return {'projectRootHash':table._hash(str(self.root).encode()),'registeredPackages':[],
                'sources':[tracks._file(p) for p in sources],'metadata':[tracks._file(p) for p in metadata]}

    def property(self,name):
        """提供明确真实属性草稿，包含实际统计目标而不继承接口名。"""
        return {'name':name,'type':'string','description':'真实属性中文说明','timing':'事件触发时','target':'gravity','remarks':''}

    def execute(self,request=None):
        """通过正式计划和一次性确认 token 执行并保存 Receipt。"""
        plan=tracks.plan_patch(self.root,request or self.request)
        plan_path='Library/Nova/AgentActions/TableSource/plan.json';receipt='Library/Nova/AgentActions/TableSource/receipt.json'
        table._atomic_write(self.root,plan_path,table._canonical(plan),state=True,replace=False)
        result=tracks.execute_plan(self.root,plan_path,plan['planId'],receipt)
        return plan,result,receipt

    def test_plan_does_not_write_source_or_output(self):
        """Plan 保持源整份 bytes 和所有目录不变，新增工作表预期摘要可重复生成。"""
        before=self.path.read_bytes();paths=set(self.root.rglob('*'))
        self.request['operations']=[{'kind':'newSheet','sheet':'打点表-Game'},
          {'kind':'appendEvent','sheet':'打点表-Game','eventName':'game_enter','chineseName':'进入游戏','description':'真实业务进入','properties':[self.property('stage')]}]
        first=tracks.plan_patch(self.root,self.request);second=tracks.plan_patch(self.root,self.request)
        self.assertEqual(first,second);self.assertEqual(before,self.path.read_bytes());self.assertEqual(paths,set(self.root.rglob('*')))

    def test_cli_registry_request_and_plan_use_fixed_cache(self):
        """CLI 从固定缓存读取登记与请求，并在同一缓存根保存计划。"""
        cache=self.root/tracks.STATE;cache.mkdir(parents=True)
        (cache/'registry.json').write_text(json.dumps(self.registry))
        (cache/'request.json').write_text(json.dumps(self.request))
        before=self.path.read_bytes()
        with contextlib.redirect_stdout(io.StringIO()) as output:
            code=tracks.main(['--project-root',str(self.root),'inspect','--source',self.request['source'],
                              '--owner-root',self.request['ownerRoot'],'--registry',tracks.STATE+'registry.json'])
        self.assertEqual(0,code);self.assertEqual('success',json.loads(output.getvalue())['status'])
        with contextlib.redirect_stdout(io.StringIO()) as output:
            code=tracks.main(['--project-root',str(self.root),'plan','--request',tracks.STATE+'request.json',
                              '--plan',tracks.STATE+'plan.json'])
        self.assertEqual(0,code);self.assertEqual(tracks.STATE+'plan.json',json.loads(output.getvalue())['plan'])
        self.assertEqual(before,self.path.read_bytes());self.assertFalse((self.root/'.agents').exists())

    def test_append_properties_preserves_targets_identifiers_and_other_parts(self):
        """在既有事件末插入属性，保留原事件/目标/说明页/无关ZIP部件并统一单色样式。"""
        original=tracks._book(self.path.read_bytes())
        self.request['operations']=[{'kind':'appendProperties','sheet':'打点表-Network','eventName':'$MPRegister','properties':[self.property('$pay_amount')]},
                                   {'kind':'setCell','sheet':'打点表-Network','eventName':'$MPRegister','propertyName':'scene','column':'属性说明','beforeValue':'当前页面','value':'当前业务页面'}]
        plan,result,receipt=self.execute();self.assertEqual('success',result['status'])
        self.assertEqual('success',tracks.verify_receipt(self.root,receipt)['status'])
        current=tracks._book(self.path.read_bytes());rows=current['sheets']['打点表-Network']['rows']
        self.assertEqual('$MPRegister',rows[1][0]);self.assertEqual('$pay_amount',rows[3][3]);self.assertEqual(['gravity','tga','gravity'],[r[7] for r in rows[1:]])
        self.assertEqual('当前业务页面',rows[2][5])
        self.assertEqual(original['parts']['custom/keep.xml'],current['parts']['custom/keep.xml'])
        self.assertEqual(original['parts']['xl/worksheets/sheet2.xml'],current['parts']['xl/worksheets/sheet2.xml'])
        book=openpyxl.load_workbook(self.path);sheet=book['打点表-Network']
        style_root=tracks.ET.fromstring(current['parts']['xl/styles.xml'])
        xfs=list(style_root.find(tracks._tag('cellXfs')));fills=list(style_root.find(tracks._tag('fills')))
        for cell in current['sheets']['打点表-Network']['xml'].iter(tracks._tag('c')):
            row=int(''.join(c for c in cell.get('r') if c.isdigit()))
            color=fills[int(xfs[int(cell.get('s'))].get('fillId'))].find(tracks._tag('patternFill')).find(tracks._tag('fgColor')).get('rgb')
            self.assertEqual('FF1F4E78' if row==1 else 'FFEAF4EC',color)
        self.assertEqual('FF1F4E78',sheet['A1'].fill.fgColor.rgb)
        self.assertEqual('宋体-简',sheet['C2'].font.name);self.assertEqual('A2:A4',str(next(m for m in sheet.merged_cells.ranges if str(m).startswith('A'))))
        book.close()
        with self.assertRaises(tracks.Error):
            tracks.execute_plan(self.root,'Library/Nova/AgentActions/TableSource/plan.json',plan['planId'],'Library/Nova/AgentActions/TableSource/again.json')

    def test_real_event_without_fixed_properties_has_no_fake_field(self):
        """真实无固定属性事件可以单行保留空属性名和空类型，不造占位字段。"""
        prop={'name':'','type':'','description':'无固定属性','timing':'事件触发时','target':'gravity','remarks':''}
        self.request['operations']=[{'kind':'appendEvent','sheet':'打点表-Network','eventName':'game_ready','chineseName':'游戏就绪','description':'真实游戏就绪时触发','properties':[prop]}]
        _,result,receipt=self.execute();self.assertEqual('success',result['status']);self.assertEqual('success',tracks.verify_receipt(self.root,receipt)['status'])
        row=tracks._book(self.path.read_bytes())['sheets']['打点表-Network']['rows'][-1]
        self.assertEqual('game_ready',row[0]);self.assertIn(row[3],(None,''));self.assertIn(row[4],(None,''));self.assertEqual('无固定属性',row[5])

    def test_no_fixed_property_cannot_mix_with_real_properties(self):
        """无固定属性说明必须真实单行，与实际属性混合的草稿在源写入前拒绝。"""
        prop={'name':'','type':'','description':'无固定属性','timing':'事件触发时','target':'gravity','remarks':''}
        self.request['operations']=[{'kind':'appendEvent','sheet':'打点表-Network','eventName':'game_ready','chineseName':'游戏就绪','description':'真实触发','properties':[prop,self.property('stage')]}]
        before=self.path.read_bytes()
        with self.assertRaises(tracks.Error): tracks.plan_patch(self.root,self.request)
        self.assertEqual(before,self.path.read_bytes())

    def test_target_change_requires_explicit_evidence(self):
        """目标不能由样式整理修改；只有定向 before/value 与真实目标证据可编辑。"""
        operation={'kind':'setCell','sheet':'打点表-Network','eventName':'$MPRegister','propertyName':'scene','column':'目标','beforeValue':'tga','value':'gravity'}
        self.request['operations']=[operation]
        with self.assertRaises(tracks.Error): tracks.plan_patch(self.root,self.request)
        operation['targetEvidence']='当前业务链已改为引力统计，用户要求该单元格同步'
        _,result,receipt=self.execute();self.assertEqual('success',result['status']);self.assertEqual('success',tracks.verify_receipt(self.root,receipt)['status'])

    def test_new_sheet_keeps_old_relationships_and_is_globally_unique(self):
        """新增页关系/内容类型闭合，全注册来源重名被拒绝，原页面内容字节不变。"""
        original=tracks._book(self.path.read_bytes())
        self.request['operations']=[{'kind':'newSheet','sheet':'打点表-Game'},
                                   {'kind':'appendEvent','sheet':'打点表-Game','eventName':'game_open','chineseName':'打开游戏','description':'业务打开事件','properties':[self.property('source')]}]
        _,result,receipt=self.execute();self.assertEqual('success',result['status']);self.assertEqual('success',tracks.verify_receipt(self.root,receipt)['status'])
        current=tracks._book(self.path.read_bytes());self.assertEqual(['打点表-Network','说明页','打点表-Game'],list(current['sheets']))
        self.assertEqual(original['parts']['xl/worksheets/sheet1.xml'],current['parts']['xl/worksheets/sheet1.xml'])
        request=copy.deepcopy(self.request);request['sourceHash']=table._hash(self.path.read_bytes());request['registry']=self.snapshot()
        with self.assertRaises(tracks.Error): tracks.plan_patch(self.root,request)

    def test_meta_source_inventory_and_code_drift_block_execute(self):
        """meta、其他源新增与实际调用代码变化都使已冻结计划在提交前阻断。"""
        plan=tracks.plan_patch(self.root,self.request);path='Library/Nova/AgentActions/TableSource/plan.json';table._atomic_write(self.root,path,table._canonical(plan),state=True)
        self.code.write_text('// 后续用户编辑\n')
        with self.assertRaises(tracks.Error): tracks.execute_plan(self.root,path,plan['planId'],'Library/Nova/AgentActions/TableSource/receipt.json')
        self.assertFalse((self.root/'Library/Nova/AgentActions/TableSource/receipt.json').exists())
        self.code.write_text('// 实际业务调用\nclass Tracking {}\n');self.path.with_name('Tracks.xlsx.meta').write_text('meta changed')
        with self.assertRaises(tracks.Error): tracks.plan_patch(self.root,self.request)

    def test_recovery_restores_only_exact_before_or_after(self):
        """精确恢复自己前后 hash，用户后续修改时拒绝覆盖且不盲删源。"""
        before=self.path.read_bytes();plan,result,receipt=self.execute();self.assertEqual('success',result['status'])
        token='restore:'+plan['planId']
        restored=tracks.recover_receipt(self.root,receipt,token)
        self.assertEqual('success',restored['status']);self.assertEqual(before,self.path.read_bytes());self.assertEqual('success',tracks.verify_receipt(self.root,receipt)['status'])

    def test_recovery_refuses_user_later_edits(self):
        """提交后任意新的整源 hash 属于用户现场，恢复不得覆盖。"""
        plan,_,receipt=self.execute()
        self.path.write_bytes(self.path.read_bytes()+b'user-later-edit')
        before=self.path.read_bytes()
        with self.assertRaises(tracks.Error): tracks.recover_receipt(self.root,receipt,'restore:'+plan['planId'])
        self.assertEqual(before,self.path.read_bytes())

    def test_cancel_after_commit_is_partial_and_not_replayable(self):
        """原子提交后中断保存 Receipt，Verify可定位，Execute计划仍只能消费一次。"""
        original=tracks._atomic_source
        def cancelled(*args):
            """模拟源替换成功后出现用户取消，不制造二次写入。"""
            original(*args);raise KeyboardInterrupt()
        with mock.patch.object(tracks,'_atomic_source',side_effect=cancelled):
            plan,result,receipt=self.execute()
        self.assertEqual('partial',result['status']);self.assertEqual('partial',tracks.verify_receipt(self.root,receipt)['status'])
        self.assertIn('completion_evidence_missing',tracks.verify_receipt(self.root,receipt)['issues'])
        self.assertEqual('success',tracks.recover_receipt(self.root,receipt,'restore:'+plan['planId'])['status'])

    def test_registered_other_workbook_and_new_sheet_duplicates_are_frozen(self):
        """新增模块源必须包含在全登记快照中，其页名冲突与整表漂移均阻断。"""
        package=self.root/'UPMPackages/com.solotopia.nova.framework.kit.game'
        path=package/'Nova/Tracks/Tracks.xlsx';path.parent.mkdir(parents=True)
        (package/'package.json').write_text(json.dumps({'name':'com.solotopia.nova.framework.kit.game','version':'1.0.0'}))
        book=openpyxl.Workbook();book.active.title='打点表-Game';book.active.append(list(tracks.HEADERS))
        book.active.append(['game_open','打开','真实打开','source','string','来源','打开时','gravity','']);book.save(path);book.close()
        with self.assertRaises(tracks.Error): tracks.plan_patch(self.root,self.request)
        request=copy.deepcopy(self.request);request['registry']=self.snapshot()
        request['operations']=[{'kind':'newSheet','sheet':'打点表-Game'}]
        with self.assertRaises(tracks.Error): tracks.plan_patch(self.root,request)
        request['source']='UPMPackages/com.solotopia.nova.framework.kit.game/Nova/Tracks/Tracks.xlsx'
        request['ownerRoot']='UPMPackages/com.solotopia.nova.framework.kit.game';request['sourceHash']=table._hash(path.read_bytes())
        request['operations']=[{'kind':'normalizeStyle','sheet':'打点表-Game'}]
        _,result,receipt=self.execute(request);self.assertEqual('success',result['status']);self.assertEqual('success',tracks.verify_receipt(self.root,receipt)['status'])
        (package/'package.json').write_text(json.dumps({'name':'com.solotopia.nova.framework.kit.game','version':'2.0.0'}))
        self.assertEqual('partial',tracks.verify_receipt(self.root,receipt)['status'])

    def test_multiple_events_insertion_keeps_original_order_and_description_segments(self):
        """前一个事件追加属性后不重排后续事件，多次结构草稿按实际行身份重新定位。"""
        self.request['operations']=[{'kind':'appendEvent','sheet':'打点表-Network','eventName':'second','chineseName':'第二事件','description':'第二事件说明','properties':[self.property('second_key')]},
                                   {'kind':'appendProperties','sheet':'打点表-Network','eventName':'$MPRegister','properties':[self.property('extra')]},
                                   {'kind':'setCell','sheet':'打点表-Network','eventName':'second','propertyName':'second_key','column':'备注','beforeValue':'','value':'第二事件原位'}]
        _,result,receipt=self.execute();self.assertEqual('success',result['status']);self.assertEqual('success',tracks.verify_receipt(self.root,receipt)['status'])
        rows=tracks._book(self.path.read_bytes())['sheets']['打点表-Network']['rows']
        self.assertEqual('$MPRegister',rows[1][0]);self.assertEqual('extra',rows[3][3]);self.assertEqual('second',rows[4][0]);self.assertEqual('第二事件原位',rows[4][8])

    def test_gateway_project_root_snapshot_roundtrip(self):
        """MCP 脱敏项目根快照可直接用于结构计划，不关闭脱敏或猜测路径。"""
        request=copy.deepcopy(self.request)
        for field in ('sources','metadata'):
            for state in request['registry'][field]:
                state['path']=state['path'].replace(str(self.root),'<project-root>')
                if state['exists'] is False:
                    state.pop('sha256')
        self.assertEqual('success',tracks.inspect_source(self.root,request['source'],request['registry'],request['ownerRoot'])['status'])
        plan,result,receipt=self.execute(request)
        self.assertEqual('success',result['status']);self.assertEqual('success',tracks.verify_receipt(self.root,receipt)['status'])
        self.assertNotIn('<project-root>',json.dumps(plan))
        request=copy.deepcopy(self.request);request['registry']['sources'][0]['path']='<project-root>/../evil.xlsx'
        with self.assertRaises(tracks.Error): tracks.plan_patch(self.root,request)

    def test_bound_coordinate_extensions_are_not_silently_moved(self):
        """数据验证和打印区域绑定含物理坐标时不做不完整移位，Plan 保留源字节。"""
        book=openpyxl.load_workbook(self.path);sheet=book['打点表-Network']
        from openpyxl.worksheet.datavalidation import DataValidation
        validation=DataValidation(type='list',formula1='"a,b"');sheet.add_data_validation(validation);validation.add('D2:D3')
        sheet.print_area='A1:I3';book.save(self.path);book.close()
        before=self.path.read_bytes();request=copy.deepcopy(self.request);request['sourceHash']=table._hash(before);request['registry']=self.snapshot()
        with self.assertRaises(tracks.Error): tracks.plan_patch(self.root,request)
        self.assertEqual(before,self.path.read_bytes())

    def test_cross_sheet_formula_is_fail_closed(self):
        """其他页引用事件坐标的公式不能因事件插入后仍保留旧范围而被当作成功。"""
        book=openpyxl.load_workbook(self.path);book['说明页']['C2']="=SUM('打点表-Network'!E2:E3)";book.save(self.path);book.close()
        before=self.path.read_bytes();request=copy.deepcopy(self.request);request['sourceHash']=table._hash(before);request['registry']=self.snapshot()
        request['operations']=[{'kind':'appendProperties','sheet':'打点表-Network','eventName':'$MPRegister','properties':[self.property('new')]}]
        with self.assertRaises(tracks.Error): tracks.plan_patch(self.root,request)
        self.assertEqual(before,self.path.read_bytes())

    def test_network_builtin_filter_database_is_updated_with_new_rows(self):
        """真实 Network 的内置筛选定义名可精确更新，不将它误判为打印或任意公式。"""
        raw=self.path.read_bytes()
        with zipfile.ZipFile(io.BytesIO(raw)) as old:
            workbook=tracks.ET.fromstring(old.read('xl/workbook.xml'))
            defined=tracks.ET.SubElement(workbook,tracks._tag('definedNames'))
            tracks.ET.SubElement(defined,tracks._tag('definedName'),name='_xlnm._FilterDatabase',localSheetId='0',hidden='1').text="'打点表-Network'!$A$1:$I$3"
            stream=io.BytesIO()
            with zipfile.ZipFile(stream,'w') as new:
                for info in old.infolist(): new.writestr(info,tracks._xml(workbook) if info.filename=='xl/workbook.xml' else old.read(info.filename))
        self.path.write_bytes(stream.getvalue())
        self.request['sourceHash']=table._hash(self.path.read_bytes());self.request['registry']=self.snapshot()
        self.request['operations']=[{'kind':'appendProperties','sheet':'打点表-Network','eventName':'$MPRegister','properties':[self.property('new')]}]
        _,result,receipt=self.execute();self.assertEqual('success',result['status']);self.assertEqual('success',tracks.verify_receipt(self.root,receipt)['status'])
        workbook=tracks._book(self.path.read_bytes())['workbook']
        actual=workbook.find(tracks._tag('definedNames')+'/'+tracks._tag('definedName')).text
        self.assertEqual("'打点表-Network'!$A$1:$I$4",actual)

    def test_same_normalize_plan_does_not_grow_styles(self):
        """已规范的事件页再整理可复用样式，不在每次操作堆叠相同样式记录。"""
        _,result,_=self.execute();self.assertEqual('success',result['status'])
        before=tracks._book(self.path.read_bytes())['parts']['xl/styles.xml']
        request=copy.deepcopy(self.request);request['sourceHash']=table._hash(self.path.read_bytes());request['registry']=self.snapshot()
        planned=tracks.plan_patch(self.root,request)
        updated,_,_=tracks._apply(tracks._book(self.path.read_bytes()),request['operations'],tracks._registry(self.root,request['registry']))
        self.assertEqual(before,tracks._book(updated)['parts']['xl/styles.xml'])
        self.assertEqual(table._hash(updated),planned['payload']['expectedHash'])

    def test_state_failure_after_plan_consumption_is_partial(self):
        """写备份失败也属于持久事务开始后的 partial，原源不变且计划不能重放。"""
        original=table._atomic_write
        def failure(root,relative,*args,**kwargs):
            """仅模拟备份状态写盘失败，其余一次性计划写入正常。"""
            if relative.endswith('-backup.json'):
                raise OSError('simulated disk failure')
            return original(root,relative,*args,**kwargs)
        before=self.path.read_bytes()
        with mock.patch.object(table,'_atomic_write',side_effect=failure):
            plan,result,receipt=self.execute()
        self.assertEqual('partial',result['status']);self.assertFalse(result['backupExists']);self.assertFalse(result['receiptExists'])
        self.assertEqual(before,self.path.read_bytes())
        with self.assertRaises(tracks.Error): tracks.execute_plan(self.root,'Library/Nova/AgentActions/TableSource/plan.json',plan['planId'],receipt)

    def test_forged_expected_hash_cannot_restore_third_party_current_bytes(self):
        """Receipt 自洽 hash 不是签名，恢复必须从备份和原草稿重算 after 才可提交。"""
        plan,_,receipt_path=self.execute()
        self.path.write_bytes(self.path.read_bytes()+b'future-user-state')
        before=self.path.read_bytes();receipt=table._read_state(self.root,receipt_path)
        original_backup=table._read_state(self.root,receipt['backup'])
        payload=receipt['plan']['payload'];payload['expectedHash']=table._hash(before)
        receipt['plan']['planId']=table._hash(table._canonical(payload))
        receipt['backup']=tracks.STATE+'tracks-'+receipt['plan']['planId']+'-backup.json'
        table._atomic_write(self.root,receipt['backup'],table._canonical(original_backup),state=True,replace=False)
        table._atomic_write(self.root,receipt_path,table._canonical(receipt),state=True)
        with self.assertRaises(tracks.Error): tracks.recover_receipt(self.root,receipt_path,'restore:'+receipt['plan']['planId'])
        self.assertEqual(before,self.path.read_bytes())

    def test_standard_printer_binary_is_preserved_and_unknown_binary_is_readonly(self):
        """实际Firebase打印设置二进制只读来源可扫描，目标标准打印部件保留，未知二进制仍拒写。"""
        with zipfile.ZipFile(self.path,'a') as archive: archive.writestr('xl/printerSettings/printerSettings1.bin',b'printer-exact')
        request=copy.deepcopy(self.request);request['sourceHash']=table._hash(self.path.read_bytes());request['registry']=self.snapshot()
        _,result,receipt=self.execute(request);self.assertEqual('success',result['status']);self.assertEqual('success',tracks.verify_receipt(self.root,receipt)['status'])
        self.assertEqual(b'printer-exact',tracks._book(self.path.read_bytes())['parts']['xl/printerSettings/printerSettings1.bin'])
        with zipfile.ZipFile(self.path,'a') as archive: archive.writestr('xl/unknown.bin',b'unknown')
        request['sourceHash']=table._hash(self.path.read_bytes());request['registry']=self.snapshot()
        before=self.path.read_bytes()
        with self.assertRaises(tracks.Error): tracks.plan_patch(self.root,request)
        self.assertEqual(before,self.path.read_bytes())

    def test_owner_and_non_event_scope_fail_closed(self):
        """任意业务目录/包缓存/厂商/说明页不因请求自称权威而得到写权限。"""
        request=copy.deepcopy(self.request);request['ownerRoot']='Assets/Game'
        with self.assertRaises(tracks.Error): tracks.plan_patch(self.root,request)
        request=copy.deepcopy(self.request);request['operations']=[{'kind':'normalizeStyle','sheet':'说明页'}]
        with self.assertRaises(tracks.Error): tracks.plan_patch(self.root,request)
        for source in ['Library/PackageCache/x/Nova/Tracks/Tracks.xlsx','../Tracks.xlsx','Assets/vendor/Tracks.xlsx']:
            with self.assertRaises(tracks.Error): tracks._path(self.root,source)


if __name__=='__main__':
    unittest.main()

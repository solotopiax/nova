#!/usr/bin/env python3
"""按冻结身份编辑项目表源；本地化结构使用专用分支，正式导出仍交对应领域 Action。"""

from __future__ import annotations

import argparse
from contextlib import contextmanager
import csv
import copy
import hashlib
import io
import json
import math
import os
from pathlib import Path, PurePosixPath
import re
import stat
import struct
import sys
import tempfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape
import zipfile


VERSION = 1
# 过程材料只读写固定本地缓存，不进入可版本管理的 Agent 配置目录。
STATE = "Library/Nova/AgentActions/TableSource/"
TYPES = {"int", "long", "float", "double", "bool", "string"}
TRACKS_HEADERS = ("事件名", "事件中文名", "事件说明", "属性名", "属性类型", "属性说明", "上报时机", "目标", "备注")
TRACKS_EDITABLE = {"事件中文名", "事件说明", "属性类型", "属性说明", "上报时机", "备注"}
# 包快照、厂商目录和正式导出物始终不是表源修改目标。
DENIED_PARTS = {
    ".git", "library", "temp", "obj", "build", "builds", "logs", "packages",
    "upmpackages", "vendor", "vendors", "thirdparty", "generated",
    "generatedlocalrepo", "jsons", "datatypes", "node_modules", ".venv",
}
PRIVATE_FIELD = re.compile(
    r"token|secret|password|passwd|credential|authorization|cookie|aes|privatekey|"
    r"apikey|appkey|clientkey|devkey|sdkkey|webhook|deviceid|openid|unionid|^uid$", re.I
)
# 属性最短匹配使自闭合空白格不会吞掉其后的非空格，样式来源必须对应精确物理格。
CELL_XML = re.compile(r"<(?P<name>(?:[\w.-]+:)?c)\b(?P<attrs>[^>]*?)(?:/\s*>|>.*?</(?P=name)\s*>)", re.S)
ROW_XML = re.compile(r"<(?P<name>(?:[\w.-]+:)?row)\b(?P<attrs>[^>]*?)(?:/\s*>|>(?P<body>.*?)</(?P=name)\s*>)", re.S)
LOCALIZATION_LAYOUT = "localization-texts"
LOCALIZATION_MAX_CELLS = 8192
TABLE_FIELDS_LAYOUT = "table-fields"
TABLE_FIELDS_MAX_CELLS = 8192


class TableSourceError(RuntimeError):
    """表示输入不唯一、源已漂移或不具备安全修改条件；错误不包含数据值。"""


def _hash(data: bytes) -> str:
    """计算整文件摘要，用于冻结输入和验证正式源写入。"""
    return hashlib.sha256(data).hexdigest()


def _canonical(value: dict) -> bytes:
    """生成严格 JSON 的稳定表示，绑定计划身份而不执行数据中的代码。"""
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


def _keys(value: dict, allowed: set, required: set, label: str) -> None:
    """拒绝未知和缺失字段，避免拼错字段导致默认扩大写入范围。"""
    if not isinstance(value, dict) or set(value) - allowed or required - set(value):
        raise TableSourceError(label + " 的字段集合无效。")


def _project_root(root: str | Path) -> Path:
    """冻结真实项目根，拒绝把符号链接目录作为本次写入根。"""
    path = Path(root).absolute()
    if any(item.is_symlink() for item in (path, *path.parents)) or not path.is_dir():
        raise TableSourceError("projectRoot 必须存在，且根与祖先目录都不能是符号链接。")
    return path.resolve(strict=True)


def _relative_path(root: Path, relative: str, *, state: bool = False, exists: bool = True) -> Path:
    """验证规范相对路径和逐层链接；状态只允许固定缓存根，普通表源仍拒绝 Library。"""
    if not isinstance(relative, str) or not relative or "\\" in relative or ":" in relative:
        raise TableSourceError("必须提供使用 / 的项目根相对路径。")
    parts = PurePosixPath(relative).parts
    if PurePosixPath(relative).is_absolute() or ".." in parts or relative != "/".join(parts):
        raise TableSourceError("路径不能是绝对路径、父目录跳转或非规范形式。")
    if state:
        if not relative.startswith(STATE) or not relative.endswith(".json"):
            raise TableSourceError("状态必须位于 Library/Nova/AgentActions/TableSource/ 的 JSON 文件。")
    elif any(part.casefold() in DENIED_PARTS for part in parts):
        raise TableSourceError("包快照、厂商目录、生成物或临时构建目录不能作为表源。")
    path = root
    for part in parts:
        path = path / part
        if path.is_symlink():
            raise TableSourceError("路径包含符号链接，不能冻结表源或状态文件。")
    if not path.resolve(strict=False).is_relative_to(root):
        raise TableSourceError("路径超出项目根。")
    if exists and not path.is_file():
        raise TableSourceError("目标文件不存在。")
    return path


def _boundaries(root: Path, source: str, source_roots: list, generated_roots: list) -> None:
    """核对调用者从有效项目配置取得的表源根与正式输出根，禁止修改生成数据。"""
    if not isinstance(source_roots, list) or not source_roots or not isinstance(generated_roots, list):
        raise TableSourceError("sourceRoots 必须非空，generatedRoots 必须显式提供。")
    source_path = root / source
    def directory(relative: str) -> Path:
        """复用路径门验证目录身份，不允许用项目根覆盖全部写入。"""
        path = _relative_path(root, relative, exists=False)
        if path.exists() and not path.is_dir():
            raise TableSourceError("源或输出边界必须是目录。")
        return path
    roots = [directory(item) for item in source_roots]
    outputs = []
    for item in generated_roots:
        # 输出可以位于 Jsons/DataTypes 等拒写目录，但仍必须留在项目且没有链接。
        if not isinstance(item, str) or not item or "\\" in item or ":" in item:
            raise TableSourceError("generatedRoots 不是规范项目相对目录。")
        parts = PurePosixPath(item).parts
        if PurePosixPath(item).is_absolute() or ".." in parts or item != "/".join(parts):
            raise TableSourceError("generatedRoots 不能超出项目。")
        target = root
        for part in parts:
            target = target / part
            if target.is_symlink():
                raise TableSourceError("generatedRoots 含符号链接。")
        outputs.append(target)
    if not any(source_path.is_relative_to(item) for item in roots):
        raise TableSourceError("目标不属于冻结的正式表源根。")
    if any(source_path.is_relative_to(item) for item in outputs):
        raise TableSourceError("目标位于正式生成输出根，不能手工编辑。")


def _private(field: str) -> bool:
    """识别凭据和身份字段，防止查询或计划公开其值。"""
    return bool(PRIVATE_FIELD.search(re.sub(r"[^a-zA-Z0-9]", "", field)))


def _type_value(value, kind: str, *, source: bool = False):
    """按明示标量类型严格验证值；源文件允许规范数字文本，不隐式转换请求类型。"""
    if not isinstance(kind, str) or kind not in TYPES:
        raise TableSourceError("仅支持 int/long/float/double/bool/string 标量，复杂类型须走所属工具。")
    if source and value is None:
        return None
    if kind == "string":
        if not isinstance(value, str):
            raise TableSourceError("string 字段必须是字符串。")
        if any((ord(char) < 32 and char not in "\t\n\r") or 0xD800 <= ord(char) <= 0xDFFF or ord(char) in (0xFFFE, 0xFFFF) for char in value):
            raise TableSourceError("字符串含不能保存到 XML 的控制字符。")
        if re.search(r"-----BEGIN .*PRIVATE KEY|^Bearer\s|^eyJ[\w-]+\.[\w-]+\.", value):
            raise TableSourceError("凭据材料不能进入表源计划或结果。")
        return value
    if source and isinstance(value, str):
        if kind in ("int", "long") and re.fullmatch(r"-?(?:0|[1-9][0-9]*)", value):
            value = int(value)
        elif kind in ("float", "double") and re.fullmatch(r"-?(?:0|[1-9][0-9]*)(?:\.[0-9]+)?(?:[eE][+-]?[0-9]+)?", value):
            value = float(value)
        elif kind == "bool" and value.lower() in ("true", "false"):
            value = value.lower() == "true"
    if kind in ("int", "long"):
        limit = 2 ** (31 if kind == "int" else 63)
        if type(value) is not int or not -limit <= value < limit:
            raise TableSourceError("整数字段类型或数值范围无效。")
    elif kind in ("float", "double"):
        if type(value) not in (int, float):
            raise TableSourceError("浮点字段必须是有限数字。")
        try:
            value = float(value)
        except OverflowError as exc:
            raise TableSourceError("浮点字段超出有限范围。") from exc
        if not math.isfinite(value):
            raise TableSourceError("浮点字段必须是有限数字。")
        if kind == "float" and abs(value) > 3.4028234663852886e38:
            raise TableSourceError("float 字段超出有限范围。")
    elif type(value) is not bool:
        raise TableSourceError("bool 字段必须是 true 或 false。")
    return value


def _layout(rows: list[list]) -> dict:
    """识别唯一 Luban ##var/##type 或普通 CSV 表头，不从第一张页猜测表结构。"""
    markers = [i for i, row in enumerate(rows) if row and row[0] == "##var"]
    if len(markers) > 1:
        raise TableSourceError("同页存在多个 ##var 表头，不能唯一定位。")
    header_index = markers[0] if markers else next((i for i, row in enumerate(rows) if any(v is not None and v != "" for v in row)), -1)
    if header_index < 0:
        raise TableSourceError("表为空，缺少字段声明。")
    header = rows[header_index]
    start = 1 if markers else 0
    fields = {}
    for index, value in enumerate(header[start:], start):
        if value is None or value == "":
            continue
        if not isinstance(value, str) or value != value.strip() or value in fields:
            raise TableSourceError("字段名必须唯一且没有首尾空白。")
        fields[value] = index
    if not fields:
        raise TableSourceError("没有可定位字段。")
    type_rows = [i for i, row in enumerate(rows) if row and row[0] == "##type"] if markers else []
    if len(type_rows) > 1 or type_rows and type_rows[0] <= header_index:
        raise TableSourceError("类型声明不唯一或位于字段声明之前。")
    types = {}
    if type_rows:
        for name, index in fields.items():
            row = rows[type_rows[0]]
            types[name] = row[index] if index < len(row) else None
    data_start = max(header_index, type_rows[0] if type_rows else header_index) + 1
    data_rows = [i for i in range(data_start, len(rows))
                 if any(v is not None and v != "" for v in rows[i])
                 and not (markers and rows[i] and isinstance(rows[i][0], str) and rows[i][0].startswith("#"))]
    return {"headerRow": header_index + 1, "fields": fields, "types": types, "dataRows": data_rows, "luban": bool(markers)}


def _csv_spans(text: str, delimiter: str) -> list[list[tuple[int, int]]]:
    """定位原 CSV 字段的文本区间，支持引号、双引号转义和字段内换行。"""
    rows, fields = [], []
    begin = index = record_begin = 0
    quoted = False
    while index < len(text):
        char = text[index]
        if char == '"':
            if quoted and index + 1 < len(text) and text[index + 1] == '"':
                index += 2
                continue
            if quoted or index == begin:
                quoted = not quoted
        if not quoted and (char == delimiter or char in "\r\n"):
            fields.append((begin, index))
            if char != delimiter:
                rows.append(fields if index > record_begin else [])
                fields = []
                if char == "\r" and index + 1 < len(text) and text[index + 1] == "\n":
                    index += 1
                record_begin = index + 1
            begin = index + 1
        index += 1
    if quoted:
        raise TableSourceError("CSV 引号未闭合。")
    if begin < len(text) or fields:
        fields.append((begin, len(text)))
        rows.append(fields)
    return rows


def _load(root: Path, relative: str, delimiter: str | None = None) -> dict:
    """读取源与声明；xlsx 仅由 openpyxl 读取，写入不重建整份工作簿。"""
    path = _relative_path(root, relative)
    raw = path.read_bytes()
    suffix = path.suffix.lower()
    result = {"source": relative, "raw": raw, "sourceHash": _hash(raw), "sheets": {}, "format": suffix[1:]}
    if suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise TableSourceError("当前 Python 缺少 openpyxl；xlsx 读取被阻塞，未安装依赖或修改表源。") from exc
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                names = archive.namelist()
                if len(names) != len(set(names)) or any("vba" in name.lower() or name.lower().endswith(".bin") for name in names):
                    raise TableSourceError("含宏/二进制扩展或重复 ZIP 成员的工作簿不在 xlsx 标量编辑范围。")
                workbook = ET.fromstring(archive.read("xl/workbook.xml"))
                relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
                targets = {r.attrib["Id"]: r.attrib["Target"] for r in relationships if r.attrib.get("TargetMode") != "External"}
                sheet_paths = {}
                for sheet in workbook.findall("{*}sheets/{*}sheet"):
                    identity = next(v for k, v in sheet.attrib.items() if k.endswith("}id"))
                    target = targets[identity]
                    sheet_paths[sheet.attrib["name"]] = target.lstrip("/") if target.startswith("/") else "xl/" + target
            book = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, read_only=False, keep_links=True)
            try:
                for sheet in book:
                    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
                    result["sheets"][sheet.title] = {"rows": rows, "part": sheet_paths[sheet.title],
                                                   "formulas": {cell.coordinate for row in sheet for cell in row if cell.data_type == "f"},
                                                   "merges": [(str(item), item.min_row, item.max_row, item.min_col - 1, item.max_col - 1)
                                                              for item in sheet.merged_cells.ranges]}
            finally:
                book.close()
        except TableSourceError:
            raise
        except Exception as exc:
            raise TableSourceError("xlsx 工作簿或关系声明无法解析。") from exc
    elif suffix == ".csv":
        try:
            text = raw.decode("utf-8-sig")
            if delimiter is None:
                candidates = [d for d in (",", ";", "\t", "|") if any(row and row[0] == "##var" and len(row) > 1 for row in csv.reader(io.StringIO(text, newline=""), delimiter=d))]
                delimiter = candidates[0] if len(candidates) == 1 else csv.Sniffer().sniff(text, delimiters=",;\t|").delimiter
            if delimiter not in (",", ";", "\t", "|"):
                raise TableSourceError("CSV delimiter 仅支持逗号、分号、Tab 或 |。")
            rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True))
            spans = _csv_spans(text, delimiter)
            if len(rows) != len(spans) or any(len(row) != len(span) for row, span in zip(rows, spans)):
                raise TableSourceError("CSV 语法不能形成可复核的字段区间。")
            result.update({"text": text, "delimiter": delimiter, "bom": raw.startswith(b"\xef\xbb\xbf")})
            result["sheets"]["csv"] = {"rows": rows, "spans": spans, "formulas": set()}
        except TableSourceError:
            raise
        except Exception as exc:
            raise TableSourceError("CSV 必须为可识别分隔符的 UTF-8 文件。") from exc
    else:
        raise TableSourceError("仅支持 .xlsx / .csv；.xls / .xlsm 与宏修改须使用专门入口。")
    return result


def _column_name(index: int) -> str:
    """把零基字段索引转换为真实 Excel 列名，地址仅作计划证据而非用户定位输入。"""
    value, result = index + 1, ""
    while value:
        value, remainder = divmod(value - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _localization_languages() -> tuple[set[str], str]:
    """从当前随包 Language 源声明读取精确枚举名与指纹，不维护第二份语言清单。"""
    source = Path(__file__).resolve().parents[2] / "Scripts/Runtime/Core/Definitions/Language.cs"
    if source.is_symlink() or not source.is_file():
        raise TableSourceError("当前 Framework 缺少真实 Language 枚举声明，不能规划新增语言。")
    raw = source.read_bytes()
    text = re.sub(r"/\*.*?\*/|//[^\r\n]*", "", raw.decode("utf-8"), flags=re.S)
    match = re.search(r"public\s+enum\s+Language\s*:\s*byte\s*\{([^{}]*)\}", text)
    if not match:
        raise TableSourceError("当前 Language 声明不在已审查的枚举语法范围。")
    names = set()
    for field in match.group(1).split(","):
        field = field.strip()
        if not field:
            continue
        entry = re.fullmatch(r"([A-Za-z][A-Za-z0-9]*)(?:\s*=\s*[0-9]+)?", field)
        if not entry or entry.group(1) in names:
            raise TableSourceError("Language 枚举项无法唯一解析。")
        names.add(entry.group(1))
    names.discard("Unspecified")
    if not names:
        raise TableSourceError("Language 枚举没有可用语言。")
    return names, _hash(raw)


def _localization_structure(sheet: dict, languages: set[str]) -> dict:
    """核对 Nova 文本四行标记、Name/Desc/语言及有效键，拒绝未知结构和公式。"""
    rows = sheet["rows"]
    if len(rows) < 5 or [rows[i][0] if rows[i] else None for i in range(4)] != ["##comment", "##var", "##type", "##comment"]:
        raise TableSourceError("本地化结构入口只支持已有 Nova 四行标记文本 Sheet。")
    if sheet["formulas"] or any(row and isinstance(row[0], str) and row[0] in ("##var", "##type") for row in rows[4:]):
        raise TableSourceError("本地化结构入口拒绝公式或重复字段/类型声明。")
    fields, language_columns = {}, {}
    for index, raw in enumerate(rows[1][1:], 1):
        if raw in (None, ""):
            continue
        if not isinstance(raw, str) or raw != raw.strip():
            raise TableSourceError("本地化表头必须使用精确名称且无首尾空白。")
        name = raw[1:] if raw.startswith("#") else raw
        if name not in {"Name", "Desc"} | languages or name in fields:
            raise TableSourceError("本地化表头含未知或重复的名称/语言。")
        if index >= len(rows[2]) or rows[2][index] != "string":
            raise TableSourceError("本地化 Name、Desc 和语言字段必须显式声明 string。")
        fields[name] = index
        if name in languages:
            language_columns[name] = index
    if "Name" not in fields or not language_columns:
        raise TableSourceError("本地化 Sheet 必须有 Name 和至少一种真实 Language。")
    name_index, keys = fields["Name"], {}
    for row_index, row in enumerate(rows[4:], 5):
        raw = row[name_index] if name_index < len(row) else None
        if raw in (None, "") or isinstance(raw, str) and raw.startswith("#"):
            continue
        key = _type_value(raw, "string")
        if key != key.strip() or not key or key in keys or _private(key):
            raise TableSourceError("本地化 Key 为空、重复、不规范或包含敏感身份。")
        if row[0] not in (None, ""):
            raise TableSourceError("本地化有效数据行不能带未知行标记。")
        for index in fields.values():
            if index < len(row) and row[index] is not None:
                _type_value(row[index], "string")
        keys[key] = row_index
    if not keys:
        raise TableSourceError("本地化结构入口要求已有有效文本行作为样式和键基线。")
    return {"fields": fields, "languages": language_columns, "keys": keys, "rowCount": len(rows),
            "merges": sheet["merges"]}


def _localization_additions(loaded: dict, request: dict) -> tuple[list[dict], dict, str]:
    """只规划新增语言空白列与追加 Key；所有新译文显式给出，不移动任何旧单元格。"""
    languages, language_hash = _localization_languages()
    if loaded["format"] != "xlsx":
        raise TableSourceError("本地化结构入口仅支持 xlsx，不支持 CSV、任意插列或删除。")
    additions, new_keys = request.get("addLanguages", []), request.get("addKeys", [])
    if not isinstance(additions, list) or not isinstance(new_keys, list) or not 1 <= len(additions) + len(new_keys) <= 256:
        raise TableSourceError("addLanguages/addKeys 必须合计 1 至 256 项。")
    structures = {name: _localization_structure(sheet, languages) for name, sheet in loaded["sheets"].items()
                  if not name.startswith("#")}
    if not structures or any(not isinstance(item, dict) or item.get("sheet") not in structures for item in additions + new_keys):
        raise TableSourceError("新增目标必须是已存在且已验证的有效本地化 Sheet。")
    sets = {tuple(sorted(item["languages"])) for item in structures.values()}
    if len(sets) != 1:
        raise TableSourceError("同工作簿有效 Sheet 的现有语言集合不一致。")
    cells, final_languages, added = [], {}, set()

    def cell(sheet_name: str, row: int, index: int, value: str, style_row: int, style_index: int) -> None:
        """冻结实际新地址、空白原值和样式来源；合并或已有内容不能成为新增写集。"""
        if row > 1048576 or index >= 16384:
            raise TableSourceError("本地化新增地址超出 Excel 行列限制。")
        sheet = loaded["sheets"][sheet_name]
        actual = sheet["rows"][row - 1][index] if row <= len(sheet["rows"]) and index < len(sheet["rows"][row - 1]) else None
        if actual not in (None, "") or any(lo <= row <= hi and left <= index <= right for _, lo, hi, left, right in sheet["merges"]):
            raise TableSourceError("本地化新增目标已有内容或位于合并区，不能覆盖或改变合并。")
        value = _type_value(value, "string")
        if value.startswith("="):
            raise TableSourceError("本地化结构填充值不能使用公式表达式。")
        cells.append({"sheet": sheet_name, "row": row, "columnIndex": index, "address": _column_name(index) + str(row),
                      "type": "string", "beforeValue": actual, "value": value,
                      "styleFrom": _column_name(style_index) + str(style_row)})

    for name, structure in structures.items():
        final_languages[name] = dict(structure["languages"])
    for item in additions:
        _keys(item, {"sheet", "language", "values"}, {"sheet", "language", "values"}, "addLanguages item")
        name, language, values = item["sheet"], item["language"], item["values"]
        if not isinstance(language, str) or language not in languages or language in final_languages[name] or (name, language) in added:
            raise TableSourceError("新增语言必须是精确的已定义 Language，且不能重复现有语言。")
        structure = structures[name]
        if not isinstance(values, list) or len(values) != len(structure["keys"]):
            raise TableSourceError("新增语言必须为每个现有有效 Key 显式提供译文，可明确填空字符串。")
        translations = {}
        for value in values:
            _keys(value, {"key", "value"}, {"key", "value"}, "language value")
            key = value["key"]
            if not isinstance(key, str) or key not in structure["keys"] or key in translations:
                raise TableSourceError("新增语言译文的 Key 未知或重复。")
            translations[key] = _type_value(value["value"], "string")
        # 只填最右字段之后的空白物理列，不调用 insert_cols，不移动原列或扩展合并。
        index = max([*structure["fields"].values(), *final_languages[name].values()]) + 1
        template = max(structure["languages"].values())
        if any(index < len(row) and row[index] not in (None, "") for row in loaded["sheets"][name]["rows"]):
            raise TableSourceError("最右字段后的待用列已有内容，不允许跳列、插列或覆盖备注。")
        cell(name, 2, index, language, 2, template)
        cell(name, 3, index, "string", 3, template)
        cell(name, 4, index, language, 4, template)
        for key, row in structure["keys"].items():
            cell(name, row, index, translations[key], row, template)
        final_languages[name][language] = index
        added.add((name, language))
    if len({tuple(sorted(item)) for item in final_languages.values()}) != 1:
        raise TableSourceError("新增语言必须覆盖同工作簿的全部有效 Sheet，最终语言集合须一致。")
    seen_keys, row_offsets = set(), {}
    for item in new_keys:
        _keys(item, {"sheet", "key", "description", "translations"}, {"sheet", "key", "description", "translations"}, "addKeys item")
        name, key = item["sheet"], _type_value(item["key"], "string")
        structure = structures[name]
        if not key or key != key.strip() or key.startswith("#") or _private(key) or key in structure["keys"] or (name, key) in seen_keys:
            raise TableSourceError("新增 Key 必须为非敏感的唯一非空名称，不能改名或覆盖旧 Key。")
        translations = item["translations"]
        if not isinstance(translations, dict) or set(translations) != set(final_languages[name]):
            raise TableSourceError("新增 Key 必须显式给出全部最终语言的译文，不自动补默认值。")
        description = _type_value(item["description"], "string")
        if "Desc" not in structure["fields"] and description:
            raise TableSourceError("当前 Sheet 没有 Desc 列，不能隐式新增描述结构。")
        row = structure["rowCount"] + row_offsets.get(name, 0) + 1
        template_row = max(structure["keys"].values())
        row_offsets[name] = row_offsets.get(name, 0) + 1
        cell(name, row, structure["fields"]["Name"], key, template_row, structure["fields"]["Name"])
        if "Desc" in structure["fields"]:
            cell(name, row, structure["fields"]["Desc"], description, template_row, structure["fields"]["Desc"])
        for language, index in final_languages[name].items():
            cell(name, row, index, translations[language], template_row, structure["languages"].get(language, max(structure["languages"].values())))
        seen_keys.add((name, key))
    if len(cells) > LOCALIZATION_MAX_CELLS:
        raise TableSourceError("本次本地化结构写集超过 8192 个单元格，请按源文件拆分明确任务。")
    proof = {name: {"fields": item["fields"], "languages": item["languages"], "keys": item["keys"],
                    "rowCount": item["rowCount"], "merges": [list(merge) for merge in item["merges"]]} for name, item in structures.items()}
    return cells, proof, language_hash


def _table_schema(root: Path, source: str, schema: str, table: str, sheet: str) -> dict:
    """只接受正式 XML 中唯一自描述表与精确 Sheet@xlsx 映射，不写外部 Schema。"""
    path = _relative_path(root, schema)
    if path.suffix.lower() != ".xml" or not all(isinstance(v, str) and v for v in (table, sheet)):
        raise TableSourceError("字段追加须明确 XML Schema、表名和 Sheet。")
    raw = path.read_bytes()
    if len(raw) > 1024 * 1024 or b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
        raise TableSourceError("字段追加 Schema 超出容量或包含外部声明。")
    try:
        document = ET.fromstring(raw)
    except ET.ParseError as exc:
        raise TableSourceError("字段追加 Schema 不是有效 XML。") from exc
    tables = [item for item in document.iter() if item.tag.rsplit("}", 1)[-1] == "table"]
    selected = [item for item in tables if item.get("name") == table]
    if len(selected) != 1:
        raise TableSourceError("字段追加必须定位唯一正式表声明。")
    item = selected[0]
    target = item.get("input", "")
    if item.get("readSchemaFromFile") != "true" or item.get("inputFiles") is not None or target.count("@") != 1:
        raise TableSourceError("仅支持 readSchemaFromFile=true 的唯一 Sheet@xlsx 自描述表。")
    selected_sheet, selected_file = target.split("@")
    mapped = _relative_path(root, str(PurePosixPath(schema).parent / selected_file))
    if selected_sheet != sheet or mapped != _relative_path(root, source) or sum(value.get("input") == target for value in tables) != 1:
        raise TableSourceError("Schema 的表、文件和 Sheet 映射不唯一或不匹配。")
    return {"schema": schema, "schemaHash": _hash(raw), "table": table, "sheet": sheet,
            "input": target, "readSchemaFromFile": True}


def _table_fields_structure(loaded: dict, sheet_name: str, row_key: str, row_key_type: str) -> dict:
    """冻结普通自描述列、唯一元数据行和全部行键；复杂列、公式或合并保持阻塞。"""
    if loaded["format"] != "xlsx" or sheet_name not in loaded["sheets"]:
        raise TableSourceError("普通字段追加只支持指定已有 xlsx Sheet。")
    sheet = loaded["sheets"][sheet_name]
    rows = sheet["rows"]
    layout = _layout(rows)
    if not layout["luban"] or sheet["formulas"] or sheet["merges"]:
        raise TableSourceError("字段追加要求自描述 Sheet，不能包含公式或合并区域。")
    markers = {name: [i + 1 for i, row in enumerate(rows) if row and row[0] == name]
               for name in ("##var", "##type", "##comment")}
    header = markers["##var"]
    types = markers["##type"]
    comments = [i for i in markers["##comment"] if header and i > header[0]]
    if len(header) != 1 or len(types) != 1 or len(comments) != 1 or not header[0] < types[0] < comments[0]:
        raise TableSourceError("字段追加要求唯一 ##var/##type/字段 ##comment，且顺序明确。")
    if any(i + 1 > comments[0] and row and isinstance(row[0], str) and row[0].startswith("##") for i, row in enumerate(rows)):
        raise TableSourceError("数据区存在额外结构指令，不能猜测 Schema。")
    fields = []
    for name, index in layout["fields"].items():
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name) or layout["types"].get(name) not in TYPES:
            raise TableSourceError("字段追加只支持普通标量列，不解释忽略列、数组、ref、enum 或复杂 Schema。")
        fields.append({"name": name, "type": layout["types"][name], "columnIndex": index,
                       "comment": rows[comments[0] - 1][index] if index < len(rows[comments[0] - 1]) else None})
    if [item["columnIndex"] for item in fields] != list(range(1, len(fields) + 1)):
        raise TableSourceError("原字段有内部空洞，不能移动旧列或猜测追加位置。")
    if not isinstance(row_key, str) or row_key not in layout["fields"] or not isinstance(row_key_type, str) or row_key_type not in {"int", "long", "string"} or layout["types"][row_key] != row_key_type or _private(row_key):
        raise TableSourceError("追加行键须为已声明的非敏感 int/long/string 字段。")
    key_column = layout["fields"][row_key]
    keys = []
    identities = set()
    for index in layout["dataRows"]:
        row = rows[index]
        if index + 1 <= comments[0] or row[0] not in (None, ""):
            raise TableSourceError("字段追加数据行带未知标记或位于元数据区。")
        raw_key = row[key_column] if key_column < len(row) else None
        key = _type_value(raw_key, row_key_type, source=True)
        if key is None or key == "" or row_key_type == "string" and (key != key.strip() or _private(key)) or key in identities:
            raise TableSourceError("字段追加的原行键为空、重复或不规范。")
        identities.add(key)
        keys.append({"key": key, "row": index + 1})
    if not keys or len(keys) > TABLE_FIELDS_MAX_CELLS:
        raise TableSourceError("字段追加要求非空且容量受限的原数据行。")
    return {"sheet": sheet_name, "rowKey": row_key, "rowKeyType": row_key_type,
            "headerRow": header[0], "typeRow": types[0], "commentRow": comments[0],
            "fields": fields, "keys": keys, "rowCount": len(rows), "merges": []}


def _table_field_specs(additions: list, proof: dict) -> list[dict]:
    """验证字段名、明示类型与 default/逐行 values 二选一，拒绝任何缺项或静默填空。"""
    if not isinstance(additions, list) or not 1 <= len(additions) <= 256:
        raise TableSourceError("字段追加须包含 1 至 256 个普通字段。")
    names = {item["name"] for item in proof["fields"]}
    keys = {item["key"] for item in proof["keys"]}
    result = []
    for item in additions:
        _keys(item, {"name", "type", "comment", "default", "values"}, {"name", "type", "comment"}, "追加字段")
        name, kind, comment = item["name"], item["type"], item["comment"]
        if not isinstance(name, str) or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name) or name in names or _private(name):
            raise TableSourceError("追加字段名必须普通、非敏感且与原字段及本批次唯一。")
        if not isinstance(kind, str) or kind not in TYPES or not isinstance(comment, str) or not comment.strip() or comment != comment.strip():
            raise TableSourceError("追加字段须有普通标量类型和明确非空注释。")
        _type_value(comment, "string")
        if ("default" in item) == ("values" in item):
            raise TableSourceError("每个追加字段须显式 default 或完整 values，不能同时提供或猜测。")
        def value(raw):
            """验证新增普通值，空值与公式样式字符串不能作为默认或行值。"""
            checked = _type_value(raw, kind)
            if checked is None or checked == "" or kind == "string" and checked.lstrip().startswith(("=", "+", "-", "@")):
                raise TableSourceError("追加字段值不能为空或包含公式样式字符串。")
            return checked
        if "default" in item:
            value(item["default"])
            result.append(dict(item))
        else:
            values = item["values"]
            if not isinstance(values, list) or len(values) != len(keys):
                raise TableSourceError("追加字段 values 必须覆盖全部原行键。")
            seen = set()
            for entry in values:
                _keys(entry, {"key", "value"}, {"key", "value"}, "追加字段行值")
                key = _type_value(entry["key"], proof["rowKeyType"])
                if key not in keys or key in seen:
                    raise TableSourceError("追加字段 values 有重复、未知或缺失行键。")
                seen.add(key)
                value(entry["value"])
            if seen != keys:
                raise TableSourceError("追加字段 values 未覆盖原行键全集。")
            result.append(dict(item))
        names.add(name)
    if len(result) * (len(keys) + 3) > TABLE_FIELDS_MAX_CELLS:
        raise TableSourceError("字段追加最多写入 8192 个单元格。")
    return result


def _table_fields_cells(proof: dict, additions: list) -> list[dict]:
    """从冻结字段顺序和原键构造唯一相邻列写集，样式只继承同行最后一个原字段。"""
    result = []
    last = max(item["columnIndex"] for item in proof["fields"])
    for offset, field in enumerate(_table_field_specs(additions, proof), 1):
        column = last + offset
        if column >= 16384:
            raise TableSourceError("字段追加超出 Excel 列界限。")
        values = ({item["key"]: item["value"] for item in field["values"]} if "values" in field else None)
        rows = [(proof["headerRow"], field["name"], "string"), (proof["typeRow"], field["type"], "string"),
                (proof["commentRow"], field["comment"], "string")]
        rows += [(item["row"], values[item["key"]] if values is not None else field["default"], field["type"]) for item in proof["keys"]]
        for row, value, kind in rows:
            result.append({"sheet": proof["sheet"], "row": row, "columnIndex": column,
                           "address": _column_name(column) + str(row), "type": kind, "beforeValue": None,
                           "value": value, "styleFrom": _column_name(last) + str(row)})
    return result


def _table_fields_preservation(loaded: dict, cells: list[dict]) -> dict:
    """为非目标 Sheet XML 与所有其他 ZIP 成员保存内容摘要，验证不能伪装旧内容仍保留。"""
    sheet = loaded["sheets"][cells[0]["sheet"]]
    targets = {item["address"] for item in cells}
    with zipfile.ZipFile(io.BytesIO(loaded["raw"])) as archive:
        text = archive.read(sheet["part"]).decode("utf-8")
        def preserved(match):
            """只略去本次新增地址的旧空格或新格，其他原单元格 XML 不作规整。"""
            identity = re.search(r'\br\s*=\s*([\'"])([^\'"]+)\1', match.group("attrs"))
            if not identity:
                raise TableSourceError("字段追加物理单元格缺少地址。")
            return "" if identity.group(2) in targets else match.group()
        text = CELL_XML.sub(preserved, text)
        dimensions = list(re.finditer(r'<(?:[\w.-]+:)?dimension\b[^>]*/\s*>', text))
        if len(dimensions) != 1:
            raise TableSourceError("字段追加要求唯一物理 dimension。")
        dimension = dimensions[0]
        text = text[:dimension.start()] + "__NOVA_TABLE_DIMENSION__" + text[dimension.end():]
        raw_members, _, _ = _zip_records(loaded["raw"], archive)
        return {"sheetPart": sheet["part"], "untouchedSheetXmlSha256": _hash(text.encode("utf-8")),
                "otherZipMembers": [{"name": item.filename, "sha256": _hash(archive.read(item.filename)),
                                     "rawRecordSha256": _hash(raw_members[item.filename])}
                                    for item in archive.infolist() if item.filename != sheet["part"]]}


def _table_fields_plan(root: Path, request: dict) -> dict:
    """以源、Schema、元数据及全行键建立严格追加计划，不把普通 Schema 编辑扩为任意代码。"""
    fields = {"source", "sourceHash", "sourceRoots", "generatedRoots", "layout", "schema", "schemaHash",
              "table", "sheet", "rowKey", "rowKeyType", "addFields"}
    _keys(request, fields, fields, "table-fields request")
    _boundaries(root, request["source"], request["sourceRoots"], request["generatedRoots"])
    loaded = _load(root, request["source"])
    schema = _table_schema(root, request["source"], request["schema"], request["table"], request["sheet"])
    if request["sourceHash"] != loaded["sourceHash"] or request["schemaHash"] != schema["schemaHash"]:
        raise TableSourceError("表源或 Schema 摘要已漂移，请重新 inspect。")
    proof = _table_fields_structure(loaded, request["sheet"], request["rowKey"], request["rowKeyType"])
    cells = _table_fields_cells(proof, request["addFields"])
    last = max(item["columnIndex"] for item in proof["fields"])
    # 新列只能填空白物理位置；包括非写入行，防止追加列成为未确认的既有内容。
    end = max(item["columnIndex"] for item in cells)
    sheet = loaded["sheets"][request["sheet"]]
    if any(value not in (None, "") for row in sheet["rows"] for value in row[last + 1:end + 1]):
        raise TableSourceError("相邻追加列已有内容，不能覆盖或移动。")
    preservation = _table_fields_preservation(loaded, cells)
    _localization_patched_bytes(loaded, cells, preserve_raw_zip=True)
    payload = {"schemaVersion": VERSION, "projectRootHash": _hash(str(root).encode()),
               "source": request["source"], "sourceHash": loaded["sourceHash"], "sourceRoots": request["sourceRoots"],
               "generatedRoots": request["generatedRoots"], "format": "xlsx", "layout": TABLE_FIELDS_LAYOUT,
               "cells": cells, "structureRequest": {key: request[key] for key in fields - {"source", "sourceHash", "sourceRoots", "generatedRoots", "layout"}},
               "structureProof": proof, "schemaProof": schema, "preservationProof": preservation}
    return {"planId": _hash(_canonical(payload)), "payload": payload, "executionStarted": False}


def _resolve_cell(loaded: dict, patch: dict) -> dict:
    """按唯一行键和字段声明定位单元格，拒绝公式、错类型和重复数据键。"""
    _keys(patch, {"sheet", "rowKey", "rowKeyType", "key", "column", "type", "beforeValue", "value"},
          {"sheet", "rowKey", "key", "column", "type", "beforeValue", "value"}, "cell patch")
    if not isinstance(patch["sheet"], str):
        raise TableSourceError("Sheet 名称必须是明示字符串。")
    sheet = loaded["sheets"].get(patch["sheet"])
    if sheet is None:
        raise TableSourceError("目标 Sheet 不存在。")
    layout = _layout(sheet["rows"])
    row_key, column = patch["rowKey"], patch["column"]
    if not isinstance(row_key, str) or not isinstance(column, str) or row_key not in layout["fields"] or column not in layout["fields"]:
        raise TableSourceError("rowKey 或 column 不是当前字段声明。")
    if row_key == column or _private(row_key) or _private(column):
        raise TableSourceError("不能修改行键或查询/写入凭据与身份字段。")
    kind = patch["type"]
    declared = layout["types"].get(column)
    key_kind = patch.get("rowKeyType", layout["types"].get(row_key))
    if declared is not None and declared != kind or layout["types"].get(row_key) is not None and key_kind != layout["types"][row_key]:
        raise TableSourceError("请求类型与当前 ##type 声明不一致。")
    wanted = _type_value(patch["key"], key_kind)
    key_index, column_index = layout["fields"][row_key], layout["fields"][column]
    seen, selected = set(), None
    for row_index in layout["dataRows"]:
        row = sheet["rows"][row_index]
        if key_index >= len(row):
            raise TableSourceError("数据行缺少行键字段。")
        coordinate = _column_name(key_index) + str(row_index + 1)
        if coordinate in sheet["formulas"]:
            raise TableSourceError("公式不能作为稳定行键。")
        key = _type_value(row[key_index], key_kind, source=True)
        if key is None or key == "" or key in seen:
            raise TableSourceError("数据行键为空、重复或类型不一致。")
        seen.add(key)
        if key == wanted:
            selected = row_index
    if selected is None:
        raise TableSourceError("目标行键不存在。")
    row = sheet["rows"][selected]
    if column_index >= len(row):
        raise TableSourceError("目标数据行缺少冻结字段。")
    coordinate = _column_name(column_index) + str(selected + 1)
    raw = row[column_index]
    if coordinate in sheet["formulas"] or loaded["format"] == "csv" and isinstance(raw, str) and raw.startswith("="):
        raise TableSourceError("不能覆盖目标公式；其他公式保持原样。")
    actual = _type_value(raw, kind, source=True)
    before = None if patch["beforeValue"] is None else _type_value(patch["beforeValue"], kind)
    after = _type_value(patch["value"], kind)
    if loaded["format"] == "csv" and kind == "string" and re.match(r"\s*[=+@-]", after):
        raise TableSourceError("CSV 字符串不能新增可能被电子表格解释为公式的内容。")
    if actual != before:
        raise TableSourceError("beforeValue 与当前源不一致，请重新 inspect/plan。")
    return {**patch, "rowKeyType": key_kind, "row": selected + 1, "columnIndex": column_index,
            "address": coordinate, "beforeValue": actual, "value": after}


def _tracks_layout(sheet: dict) -> dict:
    """核对事件九列与事件/属性唯一身份，保留 A～C 合并段而不修复未知布局。"""
    rows = sheet["rows"]
    if not rows or tuple(rows[0][:9]) != TRACKS_HEADERS or any(v not in (None, "") for row in rows for v in row[9:]):
        raise TableSourceError("Tracks 事件页必须为首行固定九列，不能套用属性说明页。")
    identities, event_rows, starts = {}, {}, {}
    current, start = None, None
    for index, row in enumerate(rows[1:], 1):
        if not any(v not in (None, "") for v in row):
            continue
        event = row[0] if row else None
        if event not in (None, ""):
            if not isinstance(event, str) or event != event.strip() or event in starts:
                raise TableSourceError("Tracks 事件名不是唯一连续分组。")
            coordinate = "A" + str(index + 1)
            if coordinate in sheet["formulas"]:
                raise TableSourceError("Tracks 事件名不能使用公式。")
            current, start = event, index + 1
            starts[event] = start
        if current is None:
            raise TableSourceError("Tracks 数据行缺少所属事件。")
        event_rows[index + 1] = current
        prop = row[3] if len(row) > 3 else None
        if prop in (None, ""):
            continue
        if not isinstance(prop, str) or prop != prop.strip() or "D" + str(index + 1) in sheet["formulas"]:
            raise TableSourceError("Tracks 属性名必须是明确文本，不能使用公式。")
        identity = (current, prop)
        if identity in identities:
            raise TableSourceError("Tracks 同事件属性名重复，不能唯一定位。")
        identities[identity] = {"selectedPropertyRow": index + 1, "eventStartRow": start}
    merges = sorted(sheet.get("merges", []))
    occupied = set()
    for _, first, last, left, right in merges:
        if first < 2 or left != right or left > 2 or first >= last or first not in event_rows:
            raise TableSourceError("Tracks 合并仅支持事件内 A～C 单列连续段。")
        for row in range(first, last + 1):
            if event_rows.get(row) != event_rows[first] or (row, left) in occupied:
                raise TableSourceError("Tracks 合并跨事件、空行或相互重叠。")
            occupied.add((row, left))
    return {"identities": identities, "merges": merges, "mergeFingerprint": _hash(_canonical({"merges": merges})),
            "eventCount": len(starts), "rowCount": len(event_rows)}


def _resolve_tracks_cell(loaded: dict, patch: dict, *, inspect: bool = False) -> dict:
    """按事件/属性复合身份定位文字叶；合并文字仅改首格，身份、目标和样式只读。"""
    fields = {"sheet", "eventName", "propertyName", "column", "type", "beforeValue", "value"}
    _keys(patch, fields, fields - {"beforeValue", "value"} if inspect else fields, "Tracks cell patch")
    if loaded["format"] != "xlsx" or any(not isinstance(patch[k], str) or not patch[k] for k in ("sheet", "eventName", "propertyName", "column")):
        raise TableSourceError("Tracks 只支持已有 xlsx 中显式页、事件、属性与列。")
    if "$" in patch["sheet"] or patch["column"] not in TRACKS_EDITABLE or patch["type"] != "string":
        raise TableSourceError("Tracks 只编辑现有事件页的文字说明，不修改身份、目标或类型布局。")
    if _private(patch["eventName"]) or _private(patch["propertyName"]):
        raise TableSourceError("Tracks 凭据或身份属性不能进入查询和计划。")
    sheet = loaded["sheets"].get(patch["sheet"])
    if sheet is None:
        raise TableSourceError("目标 Sheet 不存在。")
    layout = _tracks_layout(sheet)
    selected = layout["identities"].get((patch["eventName"], patch["propertyName"]))
    if selected is None:
        raise TableSourceError("Tracks 事件/属性身份不存在，空属性不能猜测为无固定属性。")
    column = TRACKS_HEADERS.index(patch["column"])
    row, merged_range = selected["selectedPropertyRow"], None
    for name, first, last, left, _ in layout["merges"]:
        if left == column and first <= row <= last:
            row, merged_range = first, name
            break
    coordinate = _column_name(column) + str(row)
    if coordinate in sheet["formulas"]:
        raise TableSourceError("不能覆盖 Tracks 目标公式。")
    values = sheet["rows"][row - 1]
    actual = _type_value(values[column] if column < len(values) else None, "string", source=True)
    if inspect:
        before, after = actual, actual if actual is not None else ""
    else:
        before = None if patch["beforeValue"] is None else _type_value(patch["beforeValue"], "string")
        after = _type_value(patch["value"], "string")
    if actual != before:
        raise TableSourceError("beforeValue 与当前 Tracks 源不一致，请重新 inspect/plan。")
    return {**patch, **selected, "row": row, "columnIndex": column, "address": coordinate,
            "beforeValue": actual, "value": after, "mergedRange": merged_range, "mergeFingerprint": layout["mergeFingerprint"]}


def _cell_request(cell: dict, layout: str | None) -> dict:
    """从冻结单元格重建原请求，旧标量计划与 Tracks 计划分别保留其既有输入契约。"""
    metadata = {"row", "columnIndex", "address"}
    if layout == "tracks-events":
        metadata |= {"selectedPropertyRow", "eventStartRow", "mergedRange", "mergeFingerprint"}
    return {k: value for k, value in cell.items() if k not in metadata}


def inspect_source(root: str | Path, source: str, *, sheet: str | None = None, row_key: str | None = None,
                   key=None, column: str | None = None, row_key_type: str | None = None, value_type: str | None = None,
                   delimiter: str | None = None, layout: str | None = None, event_name: str | None = None,
                   property_name: str | None = None, schema: str | None = None, table: str | None = None) -> dict:
    """只读返回实际 Sheet/字段/类型和文件摘要；定向查询才读取安全字段的原值。"""
    loaded = _load(_project_root(root), source, delimiter)
    result = {"status": "success", "source": source, "sourceHash": loaded["sourceHash"], "format": loaded["format"], "sheets": []}
    if layout == TABLE_FIELDS_LAYOUT:
        if any(v is not None for v in (key, column, value_type, delimiter, event_name, property_name)):
            raise TableSourceError("字段追加 inspect 不混用标量、Tracks 或 CSV 查询。")
        schema_proof = _table_schema(_project_root(root), source, schema, table, sheet)
        proof = _table_fields_structure(loaded, sheet, row_key, row_key_type)
        result.update({"layout": layout, "schemaHash": schema_proof["schemaHash"], "schemaProof": schema_proof,
                       "sheets": [proof], "nextFieldColumn": _column_name(len(proof["fields"]) + 1)})
        return result
    if schema is not None or table is not None:
        raise TableSourceError("schema/table 参数仅属于 table-fields 结构查询。")
    if layout == LOCALIZATION_LAYOUT:
        if loaded["format"] != "xlsx" or any(v is not None for v in (row_key, key, column, row_key_type, value_type, delimiter, event_name, property_name)):
            raise TableSourceError("本地化结构 inspect 仅接受 source、layout 和可选已有 Sheet，不混用标量查询。")
        languages, language_hash = _localization_languages()
        result.update({"layout": layout, "languageSourceHash": language_hash, "availableLanguages": sorted(languages)})
        for name, source_sheet in loaded["sheets"].items():
            if name.startswith("#") or sheet is not None and name != sheet:
                continue
            structure = _localization_structure(source_sheet, languages)
            result["sheets"].append({"name": name, "languages": list(structure["languages"]),
                                     "keys": list(structure["keys"]), "keyCount": len(structure["keys"]),
                                     "hasDescription": "Desc" in structure["fields"], "appendRow": structure["rowCount"] + 1,
                                     "nextLanguageColumn": _column_name(max(structure["fields"].values()) + 1)})
        if not result["sheets"]:
            raise TableSourceError("未找到指定的有效本地化 Sheet。")
        return result
    if layout is not None or event_name is not None or property_name is not None:
        if layout != "tracks-events" or loaded["format"] != "xlsx" or sheet not in loaded["sheets"] or any(v is not None for v in (row_key, key, row_key_type, delimiter)):
            raise TableSourceError("Tracks inspect 必须显式指定已有事件页，不混用标量行键或 CSV 入口。")
        structure = _tracks_layout(loaded["sheets"][sheet])
        result["layout"] = layout
        result["sheets"] = [{"name": sheet, "headerRow": 1, "rowCount": structure["rowCount"], "eventCount": structure["eventCount"],
                             "fields": [{"name": name, "type": "string", "editable": name in TRACKS_EDITABLE} for name in TRACKS_HEADERS]}]
        if any(v is not None for v in (event_name, property_name, column, value_type)):
            probe = _resolve_tracks_cell(loaded, {"sheet": sheet, "eventName": event_name, "propertyName": property_name,
                                                  "column": column, "type": value_type or "string"}, inspect=True)
            result["cell"] = {k: v for k, v in probe.items() if k != "value"}
        return result
    if loaded["format"] == "csv":
        result["delimiter"] = loaded["delimiter"]
    for name, table in loaded["sheets"].items():
        if sheet is not None and name != sheet:
            continue
        try:
            layout = _layout(table["rows"])
            result["sheets"].append({"name": name, "headerRow": layout["headerRow"], "rowCount": len(layout["dataRows"]),
                                      "fields": [{"name": field, "type": layout["types"].get(field), "sensitive": _private(field)} for field in layout["fields"]]})
        except TableSourceError:
            result["sheets"].append({"name": name, "status": "not_applicable", "reason": "无唯一字段声明"})
    if sheet is not None and sheet not in loaded["sheets"]:
        raise TableSourceError("目标 Sheet 不存在。")
    if any(v is not None for v in (row_key, key, column, row_key_type, value_type)):
        if sheet is None or row_key is None or key is None or column is None:
            raise TableSourceError("定向查询须提供 sheet/rowKey/key/column。")
        table, layout = loaded["sheets"][sheet], _layout(loaded["sheets"][sheet]["rows"])
        kind = value_type or layout["types"].get(column)
        # 原值从唯一安全键行获取；最终定位复用 patch 校验而非维护第二套语义。
        key_kind = row_key_type or layout["types"].get(row_key)
        if _private(row_key) or _private(column) or row_key not in layout["fields"] or column not in layout["fields"]:
            raise TableSourceError("定向查询字段无效或含凭据。")
        wanted = _type_value(key, key_kind)
        matches = [i for i in layout["dataRows"] if layout["fields"][row_key] < len(table["rows"][i])
                   and _type_value(table["rows"][i][layout["fields"][row_key]], key_kind, source=True) == wanted]
        if len(matches) != 1:
            raise TableSourceError("目标行键不存在或重复。")
        actual = _type_value(table["rows"][matches[0]][layout["fields"][column]], kind, source=True)
        probe = _resolve_cell(loaded, {"sheet": sheet, "rowKey": row_key, "rowKeyType": key_kind,
                                      "key": key, "column": column, "type": kind, "beforeValue": actual,
                                      "value": actual if actual is not None else ("" if kind == "string" else False if kind == "bool" else 0)})
        result["cell"] = {k: probe[k] for k in ("sheet", "rowKey", "rowKeyType", "key", "column", "type", "address", "beforeValue")}
    return result


def plan_patch(root: str | Path, request: dict) -> dict:
    """只读建立精确单元格计划，绑定实际源摘要、类型、原值与输入/输出边界。"""
    if isinstance(request, dict) and request.get("layout") == TABLE_FIELDS_LAYOUT:
        return _table_fields_plan(_project_root(root), request)
    if isinstance(request, dict) and request.get("layout") == LOCALIZATION_LAYOUT:
        _keys(request, {"source", "sourceHash", "sourceRoots", "generatedRoots", "layout", "addKeys", "addLanguages"},
              {"source", "sourceHash", "sourceRoots", "generatedRoots", "layout"}, "localization request")
        project = _project_root(root)
        _relative_path(project, request["source"])
        _boundaries(project, request["source"], request["sourceRoots"], request["generatedRoots"])
        loaded = _load(project, request["source"])
        if request["sourceHash"] != loaded["sourceHash"]:
            raise TableSourceError("整文件摘要已漂移，请重新 inspect。")
        cells, proof, language_hash = _localization_additions(loaded, request)
        _localization_patched_bytes(loaded, cells)
        payload = {"schemaVersion": VERSION, "projectRootHash": _hash(str(project).encode()),
                   "source": request["source"], "sourceHash": loaded["sourceHash"], "sourceRoots": request["sourceRoots"],
                   "generatedRoots": request["generatedRoots"], "format": "xlsx", "layout": LOCALIZATION_LAYOUT,
                   "cells": cells, "structureRequest": {k: request[k] for k in ("addKeys", "addLanguages") if k in request},
                   "structureProof": proof, "languageSourceHash": language_hash}
        return {"planId": _hash(_canonical(payload)), "payload": payload, "executionStarted": False}
    _keys(request, {"source", "sourceHash", "sourceRoots", "generatedRoots", "delimiter", "cells", "layout"},
          {"source", "sourceHash", "sourceRoots", "generatedRoots", "cells"}, "request")
    project = _project_root(root)
    _relative_path(project, request["source"])
    _boundaries(project, request["source"], request["sourceRoots"], request["generatedRoots"])
    loaded = _load(project, request["source"], request.get("delimiter"))
    if request["sourceHash"] != loaded["sourceHash"]:
        raise TableSourceError("整文件摘要已漂移，请重新 inspect。")
    if not isinstance(request["cells"], list) or not 1 <= len(request["cells"]) <= 256:
        raise TableSourceError("cells 必须含 1 至 256 项。")
    layout = request.get("layout")
    if "layout" in request and (layout != "tracks-events" or loaded["format"] != "xlsx" or "delimiter" in request):
        raise TableSourceError("显式 layout 只支持 xlsx tracks-events，不混用普通表或 CSV。")
    cells = [(_resolve_tracks_cell if layout == "tracks-events" else _resolve_cell)(loaded, patch) for patch in request["cells"]]
    addresses = [(cell["sheet"], cell["address"]) for cell in cells]
    if len(addresses) != len(set(addresses)):
        raise TableSourceError("同一单元格重复请求，不能合并或猜测优先级。")
    _patched_bytes(loaded, cells)
    payload = {"schemaVersion": VERSION, "projectRootHash": _hash(str(project).encode()),
               "source": request["source"], "sourceHash": loaded["sourceHash"], "sourceRoots": request["sourceRoots"],
               "generatedRoots": request["generatedRoots"], "format": loaded["format"], "cells": cells}
    if loaded["format"] == "csv":
        payload["delimiter"] = loaded["delimiter"]
    if layout is not None:
        payload["layout"] = layout
    return {"planId": _hash(_canonical(payload)), "payload": payload, "executionStarted": False}


def _cell_xml(match, value, kind: str) -> str:
    """保留目标单元格的地址/样式属性，仅替换标量类型与值节点。"""
    name, attributes = match.group("name"), match.group("attrs")
    prefix = name[:-1]
    attributes = re.sub(r"\s+t\s*=\s*(['\"]).*?\1", "", attributes)
    if kind == "string" or kind == "long" and abs(value) > 2**53:
        body = f'<{prefix}is><{prefix}t xml:space="preserve">{escape(str(value))}</{prefix}t></{prefix}is>'
        attributes += ' t="inlineStr"'
    else:
        if kind == "bool":
            attributes += ' t="b"'
        body = f'<{prefix}v>{"1" if value is True else "0" if value is False else repr(value)}</{prefix}v>'
    return f"<{name}{attributes}>{body}</{name}>"


def _zip_records(raw: bytes, archive: zipfile.ZipFile) -> tuple[dict, dict, bytes]:
    """读取标准单盘 ZIP 原始成员及目录记录；ZIP64/分盘或未知尾部明确阻塞，不静默重建。"""
    end = raw.rfind(b"PK\x05\x06", max(0, len(raw) - 65557))
    if end < 0 or end + 22 > len(raw):
        raise TableSourceError("结构追加要求标准单盘 ZIP 尾部。")
    signature, disk, directory_disk, disk_count, count, size, start, comment = struct.unpack_from("<4s4H2LH", raw, end)
    if disk or directory_disk or disk_count != count or count == 65535 or start == 0xffffffff or size == 0xffffffff or end + 22 + comment != len(raw) or start != archive.start_dir or start + size != end or count != len(archive.infolist()):
        raise TableSourceError("结构追加不解释分盘、ZIP64 或未知 ZIP 尾部。")
    ordered = sorted(archive.infolist(), key=lambda item: item.header_offset)
    members = {}
    for index, info in enumerate(ordered):
        finish = ordered[index + 1].header_offset if index + 1 < len(ordered) else start
        if not 0 <= info.header_offset < finish <= start or raw[info.header_offset:info.header_offset + 4] != b"PK\x03\x04" or info.flag_bits & 1:
            raise TableSourceError("结构追加 ZIP 成员位置或加密标记无效。")
        members[info.filename] = raw[info.header_offset:finish]
    records, position = {}, start
    for info in archive.infolist():
        if raw[position:position + 4] != b"PK\x01\x02" or position + 46 > end:
            raise TableSourceError("结构追加 ZIP 中央目录无效。")
        name_length, extra_length, comment_length = struct.unpack_from("<3H", raw, position + 28)
        finish = position + 46 + name_length + extra_length + comment_length
        if finish > end or struct.unpack_from("<L", raw, position + 42)[0] != info.header_offset:
            raise TableSourceError("结构追加 ZIP 目录与成员位置不匹配。")
        records[info.filename] = raw[position:finish]
        position = finish
    if position != end:
        raise TableSourceError("结构追加 ZIP 目录包含未知记录。")
    return members, records, raw[end:]


def _replace_raw_zip_members(raw: bytes, changes: dict[str, bytes]) -> bytes:
    """仅重编码指定 ZIP 成员；其他压缩数据、局部头、描述符及目录元数据保留，目录仅更新偏移。"""
    with zipfile.ZipFile(io.BytesIO(raw)) as original:
        members, records, end_record = _zip_records(raw, original)
        if set(changes) - set(members):
            raise TableSourceError("结构追加不能创建未知 ZIP 成员。")
        for info in original.infolist():
            if info.filename not in changes:
                continue
            stream = io.BytesIO()
            with zipfile.ZipFile(stream, "w") as target:
                target.writestr(copy.copy(info), changes[info.filename])
            encoded = stream.getvalue()
            with zipfile.ZipFile(io.BytesIO(encoded)) as generated:
                replacement_members, replacement_records, _ = _zip_records(encoded, generated)
            members[info.filename] = replacement_members[info.filename]
            records[info.filename] = replacement_records[info.filename]
        ordered = sorted(original.infolist(), key=lambda item: item.header_offset)
        result = bytearray(raw[:ordered[0].header_offset])
        offsets = {}
        for info in ordered:
            offsets[info.filename] = len(result)
            result.extend(members[info.filename])
        start = len(result)
        for info in original.infolist():
            record = bytearray(records[info.filename])
            struct.pack_into("<L", record, 42, offsets[info.filename])
            result.extend(record)
        size = len(result) - start
        tail = bytearray(end_record)
        struct.pack_into("<2L", tail, 12, size, start)
        result.extend(tail)
        return bytes(result)


def _localization_patched_bytes(loaded: dict, cells: list[dict], *, preserve_raw_zip: bool = False) -> bytes:
    """共享空白 cells/rows 追加与 dimension 更新；本地化仍全 string，Table 按冻结标量类型写入。"""
    changes = {}
    with zipfile.ZipFile(io.BytesIO(loaded["raw"])) as original:
        grouped = {}
        for cell in cells:
            grouped.setdefault(cell["sheet"], {}).setdefault(cell["row"], []).append(cell)
        for sheet_name, row_cells in grouped.items():
            part = loaded["sheets"][sheet_name]["part"]
            text = original.read(part).decode("utf-8")
            all_cells = {}
            for match in CELL_XML.finditer(text):
                identity = re.search(r'\br\s*=\s*([\'"])([^\'"]+)\1', match.group("attrs"))
                if not identity or identity.group(2) in all_cells:
                    raise TableSourceError("本地化物理单元格身份缺失或重复。")
                all_cells[identity.group(2)] = match
            rows = {}
            for match in ROW_XML.finditer(text):
                identity = re.search(r'\br\s*=\s*([\'"])([0-9]+)\1', match.group("attrs"))
                if not identity or int(identity.group(2)) in rows:
                    raise TableSourceError("本地化物理行身份缺失或重复。")
                rows[int(identity.group(2))] = match
            if not rows:
                raise TableSourceError("本地化 Sheet 缺少已有物理行。")
            replacements, new_rows = [], []
            for row_index, patches in sorted(row_cells.items()):
                match = rows.get(row_index)
                prefix = (match or next(iter(rows.values()))).group("name")[:-3]
                body = (match.group("body") or "") if match else ""
                for patch in sorted(patches, key=lambda item: item["columnIndex"]):
                    existing = next((item for item in CELL_XML.finditer(body)
                                     if re.search(r'\br\s*=\s*([\'"])' + re.escape(patch["address"]) + r'\1', item.group("attrs"))), None)
                    if existing:
                        replacement = _cell_xml(existing, patch["value"], patch["type"])
                        body = body[:existing.start()] + replacement + body[existing.end():]
                    else:
                        template = all_cells.get(patch["styleFrom"])
                        style = re.search(r'\bs\s*=\s*([\'"])([0-9]+)\1', template.group("attrs")) if template else None
                        empty = '<' + prefix + 'c r="' + patch["address"] + '"' + (' s="' + style.group(2) + '"' if style else '') + '/>'
                        replacement = _cell_xml(CELL_XML.fullmatch(empty), patch["value"], patch["type"])
                        # 新列按地址插入；所有旧节点的位置和值保持原始文本。
                        position = len(body)
                        for item in CELL_XML.finditer(body):
                            address = re.search(r'\br\s*=\s*([\'"])([A-Z]+)[0-9]+\1', item.group("attrs"))
                            if not address:
                                raise TableSourceError("本地化物理单元格地址无效。")
                            column = 0
                            for char in address.group(2):
                                column = column * 26 + ord(char) - 64
                            if column - 1 > patch["columnIndex"]:
                                position = item.start()
                                break
                        body = body[:position] + replacement + body[position:]
                if match:
                    replacement = '<' + match.group("name") + match.group("attrs") + '>' + body + '</' + match.group("name") + '>'
                    replacements.append((match.start(), match.end(), replacement))
                else:
                    # 追加行仅继承最后有效行的高度/样式，不能复制隐藏或其他结构元数据。
                    template_row = int(re.search(r'[0-9]+$', patches[0]["styleFrom"]).group())
                    template = rows.get(template_row)
                    attrs = ''.join(' ' + item.group(0) for item in re.finditer(r'\b(?:ht|customHeight|s|customFormat)\s*=\s*([\'"]).*?\1', template.group("attrs"))) if template else ''
                    new_rows.append(f'<{prefix}row r="{row_index}"{attrs}>{body}</{prefix}row>')
            for start, end, replacement in sorted(replacements, reverse=True):
                text = text[:start] + replacement + text[end:]
            if new_rows:
                closing = list(re.finditer(r'</(?:[\w.-]+:)?sheetData\s*>', text))
                if len(closing) != 1:
                    raise TableSourceError("本地化 Sheet 缺少唯一数据区。")
                position = closing[0].start()
                text = text[:position] + ''.join(new_rows) + text[position:]
            maximum_row = max(len(loaded["sheets"][sheet_name]["rows"]), max(row_cells))
            maximum_column = max(max((len(row) for row in loaded["sheets"][sheet_name]["rows"]), default=1),
                                 max(cell["columnIndex"] + 1 for patches in row_cells.values() for cell in patches))
            dimension = list(re.finditer(r'<(?:[\w.-]+:)?dimension\b[^>]*/\s*>', text))
            if len(dimension) != 1:
                raise TableSourceError("本地化 Sheet 缺少唯一 dimension，不能安全追加。")
            old = dimension[0]
            updated_dimension = re.sub(r'\bref\s*=\s*([\'"]).*?\1', 'ref="A1:' + _column_name(maximum_column - 1) + str(maximum_row) + '"', old.group())
            text = text[:old.start()] + updated_dimension + text[old.end():]
            ET.fromstring(text)
            changes[part] = text.encode("utf-8")
        if preserve_raw_zip:
            return _replace_raw_zip_members(loaded["raw"], changes)
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as target:
            target.comment = original.comment
            for info in original.infolist():
                target.writestr(info, changes.get(info.filename, original.read(info.filename)))
        return stream.getvalue()


def _patched_bytes(loaded: dict, cells: list[dict]) -> bytes:
    """生成精确新源：xlsx 仅改目标 sheet cell，CSV 仅改冻结字段文本区间。"""
    if loaded["format"] == "csv":
        replacements = []
        for cell in cells:
            if cell["beforeValue"] == cell["value"]:
                continue
            span = loaded["sheets"]["csv"]["spans"][cell["row"] - 1][cell["columnIndex"]]
            value = "true" if cell["value"] is True else "false" if cell["value"] is False else str(cell["value"])
            original = loaded["text"][span[0]:span[1]]
            if original.startswith('"') or any(c in value for c in (loaded["delimiter"], '"', "\r", "\n")):
                value = '"' + value.replace('"', '""') + '"'
            replacements.append((span, value))
        text = loaded["text"]
        for (begin, end), value in sorted(replacements, reverse=True):
            text = text[:begin] + value + text[end:]
        return (b"\xef\xbb\xbf" if loaded["bom"] else b"") + text.encode("utf-8")
    changes = {}
    with zipfile.ZipFile(io.BytesIO(loaded["raw"])) as original:
        for cell in cells:
            if cell["beforeValue"] == cell["value"]:
                continue
            part = loaded["sheets"][cell["sheet"]]["part"]
            text = changes.get(part, original.read(part).decode("utf-8"))
            matches = [m for m in CELL_XML.finditer(text) if re.search(r'\br\s*=\s*([\'"])' + re.escape(cell["address"]) + r'\1', m.group("attrs"))]
            if len(matches) != 1:
                raise TableSourceError("目标 xlsx 物理单元格缺失或重复，不能精确替换。")
            match = matches[0]
            text = text[:match.start()] + _cell_xml(match, cell["value"], cell["type"]) + text[match.end():]
            changes[part] = text
        if not changes:
            return loaded["raw"]
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as target:
            target.comment = original.comment
            for info in original.infolist():
                target.writestr(info, changes[info.filename].encode("utf-8") if info.filename in changes else original.read(info.filename))
        return stream.getvalue()


def _atomic_write(root: Path, relative: str, data: bytes, *, state: bool, replace: bool = True,
                  expected_hash: str | None = None) -> None:
    """同目录原子替换并清理临时文件；状态只写固定缓存根并保留已有文件权限。"""
    path = _relative_path(root, relative, state=state, exists=False)
    if not replace and path.exists():
        raise TableSourceError("目标状态文件已存在，不能覆盖其他计划或 Receipt。")
    path.parent.mkdir(parents=True, exist_ok=True)
    mode = stat.S_IMODE(path.stat().st_mode) if path.exists() else 0o600
    fd, temporary = tempfile.mkstemp(prefix=".nova-table-source-", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
        os.chmod(temporary, mode)
        _relative_path(root, relative, state=state, exists=False)
        if not replace and path.exists():
            raise TableSourceError("原子写入前状态文件已出现，不能覆盖其他计划或 Receipt。")
        if expected_hash is not None and (not path.is_file() or _hash(path.read_bytes()) != expected_hash):
            raise TableSourceError("原子替换前源文件已漂移，临时结果已放弃。")
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def _read_state(root: Path, relative: str) -> dict:
    """只读解析固定缓存根的状态，禁止在 JSON 中使用动态对象或代码执行。"""
    try:
        value = json.loads(_relative_path(root, relative, state=True).read_text("utf-8"))
    except (ValueError, UnicodeError) as exc:
        raise TableSourceError("状态文件不是有效 JSON。") from exc
    if not isinstance(value, dict):
        raise TableSourceError("状态根节点必须是对象。")
    return value


def _validate_table_fields_payload(root: Path, payload: dict) -> None:
    """独立校验结构收据的完整类型与派生写集，不能信任重算 planId 的任意地址或样式。"""
    request = payload["structureRequest"]
    request_fields = {"schema", "schemaHash", "table", "sheet", "rowKey", "rowKeyType", "addFields"}
    _keys(request, request_fields, request_fields, "table-fields structure request")
    _relative_path(root, request["schema"])
    proof = payload["structureProof"]
    proof_fields = {"sheet", "rowKey", "rowKeyType", "headerRow", "typeRow", "commentRow", "fields", "keys", "rowCount", "merges"}
    _keys(proof, proof_fields, proof_fields, "table-fields structure proof")
    if payload["format"] != "xlsx" or any(proof[key] != request[key] for key in ("sheet", "rowKey", "rowKeyType")):
        raise TableSourceError("字段追加计划的布局或行键身份不匹配。")
    if any(type(proof[key]) is not int for key in ("headerRow", "typeRow", "commentRow", "rowCount")) or not 1 <= proof["headerRow"] < proof["typeRow"] < proof["commentRow"] < proof["rowCount"] <= 1048576 or proof["merges"] != []:
        raise TableSourceError("字段追加计划的元数据行或合并证明无效。")
    if not isinstance(proof["fields"], list) or not proof["fields"] or not isinstance(proof["keys"], list) or not proof["keys"]:
        raise TableSourceError("字段追加计划缺少原列或原行键。")
    names = set()
    for index, field in enumerate(proof["fields"], 1):
        _keys(field, {"name", "type", "columnIndex", "comment"}, {"name", "type", "columnIndex", "comment"}, "table-fields 原字段")
        if not isinstance(field["name"], str) or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", field["name"]) or field["name"] in names or field["type"] not in TYPES or type(field["columnIndex"]) is not int or field["columnIndex"] != index:
            raise TableSourceError("字段追加计划原列身份无效。")
        names.add(field["name"])
    if proof["rowKey"] not in names or proof["rowKeyType"] not in {"int", "long", "string"} or next(field["type"] for field in proof["fields"] if field["name"] == proof["rowKey"]) != proof["rowKeyType"]:
        raise TableSourceError("字段追加计划行键声明不匹配。")
    identities = set()
    rows = set()
    for item in proof["keys"]:
        _keys(item, {"key", "row"}, {"key", "row"}, "table-fields 原行键")
        key = _type_value(item["key"], proof["rowKeyType"])
        if key in identities or type(item["row"]) is not int or not proof["commentRow"] < item["row"] <= proof["rowCount"] or item["row"] in rows:
            raise TableSourceError("字段追加计划有重复或无效原行键。")
        identities.add(key)
        rows.add(item["row"])
    if payload["cells"] != _table_fields_cells(proof, request["addFields"]):
        raise TableSourceError("字段追加计划的地址、类型、值或样式来源不是冻结结构的唯一派生写集。")
    schema = payload["schemaProof"]
    schema_fields = {"schema", "schemaHash", "table", "sheet", "input", "readSchemaFromFile"}
    _keys(schema, schema_fields, schema_fields, "table-fields Schema proof")
    if any(schema[key] != request[key] for key in ("schema", "schemaHash", "table", "sheet")) or schema["readSchemaFromFile"] is not True or not isinstance(schema["schemaHash"], str) or not re.fullmatch(r"[0-9a-f]{64}", schema["schemaHash"]):
        raise TableSourceError("字段追加计划缺少绑定的自描述 Schema。")
    preservation = payload["preservationProof"]
    _keys(preservation, {"sheetPart", "untouchedSheetXmlSha256", "otherZipMembers"}, {"sheetPart", "untouchedSheetXmlSha256", "otherZipMembers"}, "table-fields preservation proof")
    if not isinstance(preservation["untouchedSheetXmlSha256"], str) or not re.fullmatch(r"[0-9a-f]{64}", preservation["untouchedSheetXmlSha256"]) or not isinstance(preservation["otherZipMembers"], list):
        raise TableSourceError("字段追加计划原内容摘要无效。")
    for item in preservation["otherZipMembers"]:
        _keys(item, {"name", "sha256", "rawRecordSha256"}, {"name", "sha256", "rawRecordSha256"}, "table-fields ZIP proof")
        if not isinstance(item["name"], str) or any(not isinstance(item[key], str) or not re.fullmatch(r"[0-9a-f]{64}", item[key]) for key in ("sha256", "rawRecordSha256")):
            raise TableSourceError("字段追加计划 ZIP 成员摘要无效。")


def _validate_plan(root: Path, plan: dict) -> dict:
    """核对计划身份、真实项目根和严格 payload，阻止重载或手改后错误执行。"""
    _keys(plan, {"planId", "payload", "executionStarted"}, {"planId", "payload", "executionStarted"}, "plan")
    payload = plan["payload"]
    allowed = {"schemaVersion", "projectRootHash", "source", "sourceHash", "sourceRoots", "generatedRoots", "format", "delimiter", "cells", "layout"}
    if isinstance(payload, dict) and payload.get("layout") == LOCALIZATION_LAYOUT:
        allowed |= {"structureRequest", "structureProof", "languageSourceHash"}
    if isinstance(payload, dict) and payload.get("layout") == TABLE_FIELDS_LAYOUT:
        allowed |= {"structureRequest", "structureProof", "schemaProof", "preservationProof"}
    _keys(payload, allowed,
          {"schemaVersion", "projectRootHash", "source", "sourceHash", "sourceRoots", "generatedRoots", "format", "cells"}, "plan payload")
    if type(payload["schemaVersion"]) is not int or payload["schemaVersion"] != VERSION or payload["projectRootHash"] != _hash(str(root).encode()) or plan["planId"] != _hash(_canonical(payload)):
        raise TableSourceError("计划身份、版本或项目根不匹配。")
    if type(plan["executionStarted"]) is not bool:
        raise TableSourceError("计划执行状态无效。")
    _relative_path(root, payload["source"])
    _boundaries(root, payload["source"], payload["sourceRoots"], payload["generatedRoots"])
    if not isinstance(payload["sourceHash"], str) or not re.fullmatch(r"[0-9a-f]{64}", payload["sourceHash"]):
        raise TableSourceError("计划源摘要无效。")
    if payload["format"] not in ("xlsx", "csv") or Path(payload["source"]).suffix.lower() != "." + payload["format"]:
        raise TableSourceError("计划格式与正式表源不一致。")
    if payload["format"] == "csv" and payload.get("delimiter") not in (",", ";", "\t", "|") or payload["format"] == "xlsx" and "delimiter" in payload:
        raise TableSourceError("计划分隔符与源格式不一致。")
    layout = payload.get("layout")
    if layout == TABLE_FIELDS_LAYOUT:
        _keys(payload, allowed, allowed - {"delimiter"}, "table-fields plan payload")
        _validate_table_fields_payload(root, payload)
        return payload
    if layout == LOCALIZATION_LAYOUT:
        if payload["format"] != "xlsx":
            raise TableSourceError("本地化结构计划只支持 xlsx。")
        _keys(payload, allowed, allowed - {"delimiter"}, "localization plan payload")
        _keys(payload["structureRequest"], {"addKeys", "addLanguages"}, set(), "localization structure request")
        count = 0
        for name in ("addKeys", "addLanguages"):
            items = payload["structureRequest"].get(name, [])
            if not isinstance(items, list):
                raise TableSourceError("本地化结构计划请求列表无效。")
            count += len(items)
        if not 1 <= count <= 256 or not isinstance(payload["structureProof"], dict) or not payload["structureProof"]:
            raise TableSourceError("本地化结构计划缺少有效请求或原结构证明。")
        if payload["languageSourceHash"] != _localization_languages()[1]:
            raise TableSourceError("计划使用的真实 Language 声明已漂移，需重新 inspect/plan。")
        if not isinstance(payload["cells"], list) or not 1 <= len(payload["cells"]) <= LOCALIZATION_MAX_CELLS:
            raise TableSourceError("本地化结构计划写集须为 1 至 8192 个单元格。")
        addresses = []
        for cell in payload["cells"]:
            fields = {"sheet", "row", "columnIndex", "address", "type", "beforeValue", "value", "styleFrom"}
            _keys(cell, fields, fields, "localization plan cell")
            if not isinstance(cell["sheet"], str) or cell["sheet"] not in payload["structureProof"]:
                raise TableSourceError("本地化冻结 Sheet 不属于原结构证明。")
            if type(cell["row"]) is not int or not 1 <= cell["row"] <= 1048576 or type(cell["columnIndex"]) is not int or not 0 <= cell["columnIndex"] < 16384:
                raise TableSourceError("本地化冻结地址超出 Excel 界限。")
            if cell["address"] != _column_name(cell["columnIndex"]) + str(cell["row"]) or cell["type"] != "string" or cell["beforeValue"] not in (None, ""):
                raise TableSourceError("本地化新增写集只能填入已冻结空白字符串位置。")
            if not isinstance(cell["styleFrom"], str) or not re.fullmatch(r"[A-Z]+[1-9][0-9]*", cell["styleFrom"]):
                raise TableSourceError("本地化样式来源无效。")
            if _type_value(cell["value"], "string").startswith("="):
                raise TableSourceError("本地化结构计划不能写入公式。")
            addresses.append((cell["sheet"], cell["address"]))
        if len(addresses) != len(set(addresses)):
            raise TableSourceError("本地化结构计划含重复物理单元格。")
        return payload
    if "layout" in payload and (layout != "tracks-events" or payload["format"] != "xlsx"):
        raise TableSourceError("计划布局仅允许已有 xlsx Tracks 事件页。")
    if not isinstance(payload["cells"], list) or not 1 <= len(payload["cells"]) <= 256:
        raise TableSourceError("计划必须包含 1 至 256 个冻结单元格。")
    addresses = []
    for cell in payload["cells"]:
        fields = {"sheet", "column", "type", "beforeValue", "value", "row", "columnIndex", "address"}
        fields |= {"eventName", "propertyName", "selectedPropertyRow", "eventStartRow", "mergedRange", "mergeFingerprint"} if layout == "tracks-events" else {"rowKey", "rowKeyType", "key"}
        _keys(cell, fields, fields, "plan cell")
        names = ("sheet", "eventName", "propertyName", "column", "address") if layout == "tracks-events" else ("sheet", "rowKey", "column", "address")
        if any(not isinstance(cell[field], str) or not cell[field] for field in names):
            raise TableSourceError("冻结 Sheet、字段与物理地址无效。")
        if type(cell["row"]) is not int or cell["row"] < 1 or type(cell["columnIndex"]) is not int or cell["columnIndex"] < 0:
            raise TableSourceError("冻结单元格位置无效。")
        if cell["address"] != _column_name(cell["columnIndex"]) + str(cell["row"]):
            raise TableSourceError("冻结单元格地址与位置不一致。")
        if layout == "tracks-events":
            if cell["type"] != "string" or cell["column"] not in TRACKS_EDITABLE or _private(cell["eventName"]) or _private(cell["propertyName"]) or "$" in cell["sheet"]:
                raise TableSourceError("冻结 Tracks 字段不是可公开的文字叶。")
            if any(type(cell[name]) is not int or cell[name] < 2 for name in ("selectedPropertyRow", "eventStartRow")) or not cell["eventStartRow"] <= cell["row"] <= cell["selectedPropertyRow"]:
                raise TableSourceError("冻结 Tracks 事件与合并首格位置无效。")
            if cell["columnIndex"] != TRACKS_HEADERS.index(cell["column"]) or not isinstance(cell["mergeFingerprint"], str) or not re.fullmatch(r"[0-9a-f]{64}", cell["mergeFingerprint"]):
                raise TableSourceError("冻结 Tracks 列或合并摘要无效。")
            if cell["mergedRange"] is not None and (not isinstance(cell["mergedRange"], str) or not re.fullmatch(r"[ABC][1-9][0-9]*:[ABC][1-9][0-9]*", cell["mergedRange"])):
                raise TableSourceError("冻结 Tracks 合并范围无效。")
        else:
            _type_value(cell["key"], cell["rowKeyType"])
        _type_value(cell["value"], cell["type"])
        if cell["beforeValue"] is not None:
            _type_value(cell["beforeValue"], cell["type"])
        addresses.append((cell["sheet"], cell["address"]))
    if len(addresses) != len(set(addresses)):
        raise TableSourceError("计划含重复单元格。")
    return payload


def _summary(plan: dict) -> dict:
    """对外仅报告精确写入地址和类型，不输出表数据、原值或凭据。"""
    payload = plan["payload"]
    result = {"planId": plan["planId"], "source": payload["source"], "sourceHash": payload["sourceHash"],
            "cellCount": len(payload["cells"]),
            "cells": [{k: cell[k] for k in (("sheet", "eventName", "propertyName", "column", "type", "address", "mergedRange")
                                            if payload.get("layout") == "tracks-events" else ("sheet", "type", "address")
                                            if payload.get("layout") in (LOCALIZATION_LAYOUT, TABLE_FIELDS_LAYOUT) else ("sheet", "rowKey", "column", "type", "address"))} for cell in payload["cells"]]}
    if payload.get("layout") == TABLE_FIELDS_LAYOUT:
        request = payload["structureRequest"]
        result.update({"layout": TABLE_FIELDS_LAYOUT, "schema": request["schema"], "schemaHash": request["schemaHash"],
                       "table": request["table"], "sheet": request["sheet"], "rowCount": len(payload["structureProof"]["keys"]),
                       "addedFields": [{"name": item["name"], "type": item["type"],
                                        "fillMode": "default" if "default" in item else "per-key"} for item in request["addFields"]]})
    return result


@contextmanager
def _write_lock(project: Path):
    """同项目表源执行保持单写者；异常也清除本进程创建的缓存锁，不删除他人的锁。"""
    path = _relative_path(project, STATE + "write-lock.json", state=True, exists=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError as exc:
        raise TableSourceError("已有表源写入锁；确认其执行已结束后再处理，不自动删除锁。") from exc
    try:
        os.close(fd)
        yield
    finally:
        path.unlink()


def execute_plan(root: str | Path, plan_path: str, confirmation_token: str, receipt_path: str) -> dict:
    """只执行一次已确认计划；冻结原值/hash，写入前留恢复 Receipt，失败不自动回滚。"""
    project = _project_root(root)
    with _write_lock(project):
        return _execute_locked(project, plan_path, confirmation_token, receipt_path)


def _execute_locked(project: Path, plan_path: str, confirmation_token: str, receipt_path: str) -> dict:
    """在项目单写者锁内核对并消费计划；Receipt 与源写入保持明确的恢复边界。"""
    if plan_path == receipt_path:
        raise TableSourceError("Plan 与 Receipt 必须使用不同文件。")
    plan = _read_state(project, plan_path)
    payload = _validate_plan(project, plan)
    if plan["executionStarted"] or confirmation_token != plan["planId"]:
        raise TableSourceError("计划已消费或缺少绑定计划的确认 token。")
    loaded = _load(project, payload["source"], payload.get("delimiter"))
    if loaded["sourceHash"] != payload["sourceHash"]:
        raise TableSourceError("源文件已漂移，未执行计划。")
    layout = payload.get("layout")
    if layout in (LOCALIZATION_LAYOUT, TABLE_FIELDS_LAYOUT):
        request = {k: payload[k] for k in ("source", "sourceHash", "sourceRoots", "generatedRoots", "layout")}
        request.update(payload["structureRequest"])
        reconstructed = plan_patch(project, request)
        if reconstructed["payload"] != payload:
            raise TableSourceError("新增写集、样式来源、Schema 或原结构已漂移。")
        updated = _localization_patched_bytes(loaded, payload["cells"], preserve_raw_zip=layout == TABLE_FIELDS_LAYOUT)
    else:
        cells = [(_resolve_tracks_cell if layout == "tracks-events" else _resolve_cell)(loaded, _cell_request(cell, layout)) for cell in payload["cells"]]
        if cells != payload["cells"]:
            raise TableSourceError("真实行键、字段或类型位置已漂移。")
        updated = _patched_bytes(loaded, cells)
    receipt = {"schemaVersion": VERSION, "plan": {**plan, "executionStarted": True},
               "expectedFileHash": _hash(updated), "completed": False}
    _atomic_write(project, receipt_path, _canonical(receipt) + b"\n", state=True, replace=False)
    plan["executionStarted"] = True
    _atomic_write(project, plan_path, _canonical(plan) + b"\n", state=True)
    try:
        if _hash(_relative_path(project, payload["source"]).read_bytes()) != payload["sourceHash"]:
            raise TableSourceError("写入前源文件已漂移。")
        if layout == TABLE_FIELDS_LAYOUT and _table_schema(project, payload["source"], payload["structureRequest"]["schema"], payload["structureRequest"]["table"], payload["structureRequest"]["sheet"]) != payload["schemaProof"]:
            raise TableSourceError("写入前 Schema 已漂移。")
        if updated != loaded["raw"]:
            _atomic_write(project, payload["source"], updated, state=False, expected_hash=payload["sourceHash"])
        receipt["completed"] = True
        _atomic_write(project, receipt_path, _canonical(receipt) + b"\n", state=True)
    except (OSError, TableSourceError):
        return {"status": "partial", **_summary(plan), "receipt": receipt_path,
                "reason": "写入开始后未取得完整完成证据；保留现场，仅使用 Verify，不重放 Execute。"}
    return {"status": "success", **_summary(plan), "receipt": receipt_path,
            "changed": updated != loaded["raw"], "fileHash": receipt["expectedFileHash"]}


def verify_receipt(root: str | Path, receipt_path: str) -> dict:
    """只读核对当前整文件hash和冻结目标值，不再次写入、导出或修复表源。"""
    project = _project_root(root)
    receipt = _read_state(project, receipt_path)
    _keys(receipt, {"schemaVersion", "plan", "expectedFileHash", "completed"},
          {"schemaVersion", "plan", "expectedFileHash", "completed"}, "receipt")
    payload = _validate_plan(project, receipt["plan"])
    if type(receipt["schemaVersion"]) is not int or receipt["schemaVersion"] != VERSION or type(receipt["completed"]) is not bool:
        raise TableSourceError("Receipt 版本或完成标记无效。")
    if not receipt["plan"]["executionStarted"] or not isinstance(receipt["expectedFileHash"], str) or not re.fullmatch(r"[0-9a-f]{64}", receipt["expectedFileHash"]):
        raise TableSourceError("Receipt 必须属于已消费计划，且具有有效整文件摘要。")
    loaded = _load(project, payload["source"], payload.get("delimiter"))
    issues = []
    if loaded["sourceHash"] != receipt["expectedFileHash"]:
        issues.append("file_hash_changed")
    for cell in payload["cells"]:
        layout = payload.get("layout")
        if layout in (LOCALIZATION_LAYOUT, TABLE_FIELDS_LAYOUT):
            sheet = loaded["sheets"].get(cell["sheet"])
            if sheet is None or cell["address"] in sheet["formulas"] or cell["row"] > len(sheet["rows"]) or cell["columnIndex"] >= len(sheet["rows"][cell["row"] - 1]):
                issues.append("cell_value_or_key_changed")
            else:
                actual = sheet["rows"][cell["row"] - 1][cell["columnIndex"]]
                try:
                    observed = _type_value(actual, cell["type"], source=True) if layout == TABLE_FIELDS_LAYOUT else (actual if actual is not None else "")
                except TableSourceError:
                    observed = None
                if observed != cell["value"]:
                    issues.append("cell_value_or_key_changed")
            continue
        probe = _cell_request(cell, layout)
        probe["beforeValue"] = probe["value"]
        try:
            observed = (_resolve_tracks_cell if layout == "tracks-events" else _resolve_cell)(loaded, probe)
            if any(observed[name] != cell[name] for name in (("address", "selectedPropertyRow", "eventStartRow", "mergedRange", "mergeFingerprint") if layout == "tracks-events" else ("address",))):
                issues.append("cell_position_changed")
        except TableSourceError:
            issues.append("cell_value_or_key_changed")
    if payload.get("layout") == TABLE_FIELDS_LAYOUT:
        try:
            request, proof = payload["structureRequest"], payload["structureProof"]
            observed = _table_fields_structure(loaded, request["sheet"], request["rowKey"], request["rowKeyType"])
            expected_fields = proof["fields"] + [{"name": field["name"], "type": field["type"], "columnIndex": len(proof["fields"]) + index, "comment": field["comment"]}
                                               for index, field in enumerate(request["addFields"], 1)]
            if observed["fields"] != expected_fields or any(observed[key] != proof[key] for key in ("headerRow", "typeRow", "commentRow", "keys", "rowCount", "merges")):
                issues.append("table_fields_structure_changed")
            if _table_schema(project, payload["source"], request["schema"], request["table"], request["sheet"]) != payload["schemaProof"]:
                issues.append("table_schema_changed")
            if _table_fields_preservation(loaded, payload["cells"]) != payload["preservationProof"]:
                issues.append("table_original_content_changed")
        except (TableSourceError, KeyError, TypeError):
            issues.append("table_fields_structure_changed")
    if payload.get("layout") == LOCALIZATION_LAYOUT:
        try:
            languages = _localization_languages()[0]
            structures = {name: _localization_structure(loaded["sheets"][name], languages) for name in payload["structureProof"]}
            if len({tuple(sorted(item["languages"])) for item in structures.values()}) != 1:
                issues.append("localization_structure_changed")
            for name, proof in payload["structureProof"].items():
                observed = structures[name]
                if any(observed["fields"].get(field) != index for field, index in proof["fields"].items()) or any(observed["keys"].get(key) != row for key, row in proof["keys"].items()) or [list(merge) for merge in observed["merges"]] != proof["merges"]:
                    issues.append("localization_structure_changed")
        except (TableSourceError, KeyError, TypeError):
            issues.append("localization_structure_changed")
    if not receipt["completed"]:
        issues.append("completion_evidence_missing")
    return {"status": "partial" if issues else "success", **_summary(receipt["plan"]),
            "receipt": receipt_path, "fileHash": loaded["sourceHash"], "issues": sorted(set(issues))}


def main(argv: list[str] | None = None) -> int:
    """提供固定 CLI；请求可读固定缓存根，所有输出结构化且错误不回显输入数据。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", required=True)
    commands = parser.add_subparsers(dest="operation", required=True)
    inspect = commands.add_parser("inspect")
    inspect.add_argument("--source", required=True)
    for name in ("sheet", "row-key", "column", "row-key-type", "type", "delimiter", "layout", "event-name", "property-name", "schema", "table"):
        inspect.add_argument("--" + name)
    inspect.add_argument("--key-json")
    plan = commands.add_parser("plan")
    plan.add_argument("--request", required=True, help="项目内请求 JSON；- 表示 stdin")
    plan.add_argument("--plan", required=True)
    execute = commands.add_parser("execute")
    execute.add_argument("--plan", required=True)
    execute.add_argument("--confirmation-token", required=True)
    execute.add_argument("--receipt", required=True)
    verify = commands.add_parser("verify")
    verify.add_argument("--receipt", required=True)
    args = parser.parse_args(argv)
    try:
        project = _project_root(args.project_root)
        if args.operation == "inspect":
            result = inspect_source(project, args.source, sheet=args.sheet, row_key=args.row_key,
                                    key=json.loads(args.key_json) if args.key_json is not None else None,
                                    column=args.column, row_key_type=args.row_key_type, value_type=args.type, delimiter=args.delimiter,
                                    layout=args.layout, event_name=args.event_name, property_name=args.property_name,
                                    schema=args.schema, table=args.table)
        elif args.operation == "plan":
            if args.request == "-":
                request = json.loads(sys.stdin.read())
            elif args.request.startswith(STATE):
                request = _read_state(project, args.request)
            else:
                request = json.loads(_relative_path(project, args.request).read_text("utf-8"))
            planned = plan_patch(project, request)
            _atomic_write(project, args.plan, _canonical(planned) + b"\n", state=True, replace=False)
            result = {"status": "success", **_summary(planned), "plan": args.plan, "executionStarted": False}
        elif args.operation == "execute":
            result = execute_plan(project, args.plan, args.confirmation_token, args.receipt)
        else:
            result = verify_receipt(project, args.receipt)
    except TableSourceError as exc:
        result = {"status": "blocked", "reason": str(exc)}
    except (OSError, ValueError, TypeError, KeyError):
        result = {"status": "blocked", "reason": "文件、JSON 或输入类型无法形成安全的表源操作；未回显输入。"}
    print(json.dumps(result, ensure_ascii=False, allow_nan=False))
    return 0 if result["status"] == "success" else 3 if result["status"] == "partial" else 2


if __name__ == "__main__":
    raise SystemExit(main())
